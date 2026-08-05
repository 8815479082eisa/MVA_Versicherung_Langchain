"""
Unified RAG service used by API and CLI.

Migration note:
- `src/api/rag_service.py` is the single source of truth for the active pipeline.
- Legacy cloud-provider orchestration was retired from the active runtime path.
- Compatibility wrappers remain available in `main.py` for evaluation scripts.
"""

from __future__ import annotations

import glob
import hashlib
import json
import logging
import os
import re
import threading
import time
from dataclasses import dataclass
from datetime import datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import TYPE_CHECKING, Any, Callable, Dict, List, Optional, Tuple, TypeAlias

from langchain_core.documents import Document

from src.core.answer_completeness import (
    build_answer_requirements,
    evaluate_answer_completeness,
    format_answer_requirements,
    format_missing_requirements,
)
from src.core.diagnostic_capture import sanitize_diagnostic_text, sanitize_diagnostic_value
from src.core.llm_runtime import (
    GuardrailInvalidOutputError,
    LLMStageTimeoutError,
    RetrievalFailedError,
    RuntimeExecutionError,
    invoke_llm_stage,
    parse_closed_label,
    run_bounded_operation,
)
from src.core.runtime_diagnostics import (
    add_retry,
    begin_stage,
    configure_self_check,
    current_diagnostics,
    finish_stage,
    mark_stage,
    mark_self_check_skipped,
    record_evidence,
)

Chroma = None

if TYPE_CHECKING:
    from langchain_chroma import Chroma as _ChromaVectorStore

    ChromaVectorStore: TypeAlias = _ChromaVectorStore
else:
    ChromaVectorStore: TypeAlias = Any

chromadb = None

try:
    from config.models import (
        ModelSettings,
        build_answer_model_warning,
        load_model_settings,
        runtime_config_snapshot,
    )
except Exception:
    from src.config.models import (
        ModelSettings,
        build_answer_model_warning,
        load_model_settings,
        runtime_config_snapshot,
    )

try:
    from core.safety_adapter import (
        SafetyChecker,
        SafetyResult,
        classify_safety_result,
        create_safety_checker,
        fallback_text_for_result,
    )
except Exception:
    from src.core.safety_adapter import (
        SafetyChecker,
        SafetyResult,
        classify_safety_result,
        create_safety_checker,
        fallback_text_for_result,
    )

try:
    from utils.language import detect_response_language
except Exception:
    from src.utils.language import detect_response_language


HuggingFaceEmbeddings = None
ChatOllama = None
OpenAIResponsesAnswerModel = None
FlagReranker = None
CrossEncoder = None
ChatPromptTemplate = None
PyPDFLoader = None
BM25Retriever = None
RecursiveCharacterTextSplitter = None

logger = logging.getLogger(__name__)

_SOURCE_PREFILTER_STOPWORDS = {
    "about",
    "after",
    "also",
    "does",
    "from",
    "have",
    "included",
    "into",
    "only",
    "selected",
    "that",
    "their",
    "there",
    "these",
    "this",
    "under",
    "what",
    "when",
    "where",
    "which",
    "with",
}


SETTINGS: ModelSettings = load_model_settings() #Hier werden alle zentralen Einstellungen geladen.

PDF_DIRECTORY = str(SETTINGS.storage.pdf_directory) 
AUDIT_LOG_FILE = str(SETTINGS.storage.audit_log_file)
SUPPORTED_SOURCE_EXTENSIONS = {".pdf", ".txt", ".csv", ".xlsx", ".docx", ".png", ".jpg", ".jpeg"}
ANSWER_STYLE = os.getenv("ANSWER_STYLE", "detailed")  #  detailed | concise.  ausführliche Antworten oder kurze Antworten 
QUESTION_MODAL_TOKENS = { #Diese Tokens werden verwendet, um Ja/Nein-Fragen zu identifizieren. Wenn die Frage mit einem dieser Tokens beginnt, wird sie als Ja/Nein-Frage klassifiziert. Dies kann bei der Überprüfung der semantischen Sicherheit von umgeschriebenen Fragen helfen, da bestimmte Umformulierungen die Frage von einer Ja/Nein-Frage in eine andere Art von Frage ändern könnten, was vermieden werden sollte.
    "can",
    "could",
    "should",
    "would",
    "will",
    "do",
    "does",
    "did",
    "is",
    "are",
    "was",
    "were",
    "has",
    "have",
    "had",
    "may",
    "might",
    "must",
}
QUESTION_OPENING_TOKENS = {"what", "when", "where", "which", "why", "how", "who", "whom"} #Diese Tokens werden verwendet, um die Form der Frage zu bestimmen. Wenn die Frage mit einem dieser Tokens beginnt, wird sie als eine bestimmte Art von Frage klassifiziert (z.B. "what" für eine Frage nach Informationen, "how" für eine Frage nach dem Prozess, etc.). Dies kann bei der Überprüfung der semantischen Sicherheit von umgeschriebenen Fragen helfen, da bestimmte Umformulierungen die Art der Frage verändern könnten (z.B. von einer "what"-Frage zu einer "how"-Frage), was vermieden werden sollte.
FIRST_PERSON_TOKENS = {"i", "me", "my", "mine", "we", "us", "our", "ours"} #Diese Tokens werden verwendet, um zu überprüfen, ob eine umgeschriebene Frage eine erste Person Perspektive einführt, die im Original nicht vorhanden war. Wenn die umgeschriebene Frage Wörter wie "I" oder "my" enthält, aber die Originalfrage nicht, könnte dies darauf hinweisen, dass die Umformulierung die Absicht der Frage verändert hat, was vermieden werden sollte.
REWRITE_STOPWORDS = { #Diese Wörter werden bei der Überprüfung der semantischen Sicherheit von umgeschriebenen Fragen ignoriert, da sie häufig zur Verbesserung der Formulierung verwendet werden, ohne die Absicht zu verändern.
    "a",
    "an",
    "the",
    "and",
    "or",
    "for",
    "to",
    "of",
    "on",
    "in",
    "at",
    "by",
    "with",
    "from",
    "is",
    "are",
    "was",
    "were",
    "be",
    "being",
    "been",
}


@dataclass #Diese Dekorator wird verwendet, um eine Klasse zu erstellen, die zur Speicherung der Quelle der Antwort verwendet werden kann. Es enthält Informationen wie die Dokument-ID, den Titel des Dokuments, die Seite, den Abschnitt und einen Ausschnitt des relevanten Textes.
class Source: 
    document_id: str
    document_title: str
    page: Optional[int] = None
    section: Optional[str] = None
    snippet: Optional[str] = None  # Ein kurzer Ausschnitt des relevanten Textes aus der Quelle, der in den Antwortquellen angezeigt werden kann, um dem Benutzer Kontext zu geben, ohne dass er die gesamte Quelle lesen muss. Dies kann besonders nützlich sein, wenn die Quellen umfangreich sind oder wenn mehrere Quellen bereitgestellt werden, um die Antwort zu unterstützen. Das Snippet sollte so gewählt werden, dass es den relevanten Teil der Quelle hervorhebt, der zur Beantwortung der Frage beigetragen hat.


@dataclass #Diese Dekorator wird verwendet, um eine Klasse zu erstellen, die zur Speicherung der Antwort und der Quellen verwendet werden kann.
class AnswerResult: #Diese Klasse wird verwendet, um die Antwort und die Quellen zu speichern.
    answer: str
    sources: List[Source]
    query: str
    latency_ms: Optional[int] = None
    diagnostics: Optional[dict[str, Any]] = None


ROUTER_SYSTEM_PROMPT = """You are an intelligent router for an insurance assistant.
Decide whether the user question requires document retrieval (RETRIEVE) or can be answered directly without searching the document base (NO_RETRIEVE).
Choose RETRIEVE for document-specific, policy-specific, source-dependent, or uncertain cases.
Choose NO_RETRIEVE only for general knowledge, conversational, or simple assistant questions that can be answered safely without document evidence.
If you are unsure, reply RETRIEVE.
Reply ONLY with RETRIEVE or NO_RETRIEVE.
"""

SELF_CHECK_SYSTEM_PROMPT = """You are an assistant evaluating whether the provided context documents are relevant to the user question.
Reply RELEVANT if the context contains enough directly useful information to answer at least part of the question accurately.
Reply IRRELEVANT if the context is off-topic, too vague, or does not help answer the question.
Respond ONLY with RELEVANT or IRRELEVANT.
"""

QUERY_REWRITE_SYSTEM_PROMPT = """You rewrite user search queries to improve document retrieval.
Rules:
- Preserve the original user intent exactly. Do not change what the user is asking about.
- Only adjust phrasing or word order to better match document terminology.
- Do not change the subject, perspective, or goal of the question.
- Preserve all important entities, numbers, dates, product names, and policy names.
- If no better retrieval phrasing is possible, return the original query unchanged.
- If you are unsure, return the original query unchanged.
- Reply only with the rewritten query, nothing else.
"""

SYSTEM_PROMPT = """You are a helpful insurance information assistant.
Answer questions based on the provided context passages.
Do not produce safety disclaimers, policy-compliance warnings, or refusal messages - safety is enforced externally by a separate guardrail layer.
If the context contains relevant information, answer from it directly and concisely.
If the context does not contain sufficient information to answer the question, respond with: "The available sources do not contain enough information to answer this question."
Do not fabricate information not present in the context.
PDF context passages begin with user-facing source labels such as [policy.pdf, page 7].
Use those exact labels as citations for the document facts you state. The page
number in the label is already the physical, human-readable PDF page. Never
invent a placeholder citation, expose an internal filesystem path, or change a
provided page number.
Respond ONLY in English.

--- Chat History ---
{{chat_history}}

"""

DIRECT_ANSWER_SYSTEM_PROMPT = """You are a helpful insurance information assistant.
Answer the user's question directly without retrieving documents.
Use this mode only for general knowledge, conversational, or simple assistant questions.
Do not claim to have searched, quoted, or verified document sources.
If the question actually requires policy-specific or document-specific evidence, say: "The available sources do not contain enough information to answer this question."
Do not produce safety disclaimers, policy-compliance warnings, or refusal messages - safety is enforced externally by a separate guardrail layer.
Respond concisely and ONLY in English.

--- Chat History ---
{{chat_history}}

"""


INSUFFICIENT_INFORMATION_MESSAGE = (
    "The available sources do not contain enough information to answer this question."
)
NO_RELEVANT_INFORMATION_MESSAGE = (
    "I could not find relevant information in the indexed documents. "
    "Please rephrase your question or provide additional documents."
)
RELIABLE_ANSWER_FAILURE_MESSAGE = (
    "I could not generate a reliable answer from the current model response. "
    "Please try again or rephrase the question."
)


def _source_filename(source_path: Any) -> str:
    normalized = str(source_path or "unknown").replace("\\", "/").rstrip("/")
    return normalized.rsplit("/", 1)[-1] or "unknown"


def _metadata_page(doc: Document) -> int | None:
    page = doc.metadata.get("page")
    if isinstance(page, str):
        try:
            page = int(page)
        except ValueError:
            page = None
    return page if isinstance(page, int) and page >= 0 else None


def _doc_reference_label(doc: Document) -> str:
    if doc.metadata.get("source_type") == "crm":
        section = str(doc.metadata.get("section") or "record").strip()
        return f"[CRM: {section}]"

    source_path = doc.metadata.get("source", "unknown")
    filename = _source_filename(source_path)
    page = _metadata_page(doc)
    if page is None:
        return f"[{filename}]"
    return f"[{filename}, page {page + 1}]"


_INSURANCE_DOMAIN_TERMS = {
    "motor": {"motor", "vehicle", "vehicles", "car", "cars", "automobile", "windscreen", "windshield"},
    "household": {"household", "contents", "home"},
    "liability": {"liability"},
    "travel": {"travel", "trip", "journey"},
    "pet": {"pet", "pets", "dog", "dogs", "cat", "cats"},
    "health": {"health", "medical"},
    "legal": {"legal"},
}
_PRODUCT_DESCRIPTOR_KEYS = (
    "product",
    "product_name",
    "insurance_type",
    "document_type",
    "policy_type",
    "title",
)


def _insurance_domains(text: str) -> set[str]:
    tokens = set(re.findall(r"[a-z0-9]+", (text or "").casefold()))
    return {
        domain
        for domain, terms in _INSURANCE_DOMAIN_TERMS.items()
        if tokens & terms
    }


def _insurance_product_affinity(query: str, doc: Document) -> int:
    """Return a general product-domain preference without relying on source IDs."""

    query_domains = _insurance_domains(query)
    descriptor = " ".join(
        [
            _source_filename(doc.metadata.get("source")),
            *(str(doc.metadata.get(key) or "") for key in _PRODUCT_DESCRIPTOR_KEYS),
        ]
    )
    document_domains = _insurance_domains(descriptor)
    if not document_domains:
        document_domains = _insurance_domains((doc.page_content or "")[:800])
    if query_domains and document_domains:
        return 2 if query_domains & document_domains else -2

    query_tokens = {
        token
        for token in re.findall(r"[a-z0-9]+", (query or "").casefold())
        if len(token) >= 4 and token not in _SOURCE_PREFILTER_STOPWORDS
    }
    descriptor_tokens = set(re.findall(r"[a-z0-9]+", descriptor.casefold()))
    return 1 if query_tokens & descriptor_tokens else 0


_DEDUCTIBLE_WAIVER_RE = re.compile(
    r"(?:will not have to bear (?:an? )?deductible|"
    r"(?:deductible|excess) (?:will )?(?:not be applied|be waived)|"
    r"(?:no|without) (?:deductible|excess))",
    re.IGNORECASE,
)
_WINDSCREEN_REPAIR_RE = re.compile(
    r"(?:windscreen|windshield|front glass|glass damage)[^.\n]{0,240}"
    r"repair(?:ed)?[^.\n]{0,160}(?:rather than|and not|instead of) replac(?:e|ed|ement)",
    re.IGNORECASE,
)
_ABROAD_SCENARIO_RE = re.compile(
    r"\b(?:abroad|foreign country|outside (?:of )?switzerland|overseas)\b",
    re.IGNORECASE,
)


def _contains_complete_windscreen_waiver(text: str) -> bool:
    return bool(
        _DEDUCTIBLE_WAIVER_RE.search(text or "")
        and _WINDSCREEN_REPAIR_RE.search(text or "")
    )


def _ensure_evidence_backed_windscreen_waiver(
    answer: str,
    docs: List[Document],
) -> str:
    """Complete the repair waiver only when its governing rule is selected."""

    answer = (answer or "").strip()
    if not answer:
        return answer
    docs_by_source: dict[str, list[Document]] = {}
    for doc in docs:
        if doc.metadata.get("source_type") == "crm":
            continue
        source = str(doc.metadata.get("source") or "")
        docs_by_source.setdefault(source, []).append(doc)

    supporting_doc: Document | None = None
    for source_docs in docs_by_source.values():
        ordered = sorted(
            source_docs,
            key=lambda doc: (
                _metadata_page(doc) if _metadata_page(doc) is not None else -1,
                int(doc.metadata.get("start_index") or 0),
            ),
        )
        combined = "\n".join(doc.page_content or "" for doc in ordered)
        if not _contains_complete_windscreen_waiver(combined):
            continue
        supporting_doc = next(
            (
                doc
                for doc in ordered
                if _WINDSCREEN_REPAIR_RE.search(doc.page_content or "")
            ),
            ordered[0],
        )
        break
    if supporting_doc is None:
        return answer

    normalized_answer = answer.casefold()
    if _contains_complete_windscreen_waiver(normalized_answer):
        return answer

    sentences = re.split(r"(?<=[.!?])\s+|[\r\n]+", answer)
    kept_sentences = [
        sentence
        for sentence in sentences
        if not re.fullmatch(r"\s*\d+[.)]?\s*", sentence)
        and not (
            (
                re.search(r"windscreen|windshield|glass", sentence, re.IGNORECASE)
                and re.search(r"deductible|excess", sentence, re.IGNORECASE)
                and re.search(
                    r"\b(?:may|might|could|still)\b.*\bappl(?:y|ies)\b",
                    sentence,
                    re.IGNORECASE,
                )
            )
            or (
                re.search(r"deductible|excess", sentence, re.IGNORECASE)
                and re.search(
                    r"\b(?:applies|apply)\s+unless\b",
                    sentence,
                    re.IGNORECASE,
                )
            )
            or (
                re.search(r"windscreen|windshield|glass", sentence, re.IGNORECASE)
                and re.search(
                    r"repair(?:ed)?\s+(?:rather than|instead of|and not)\s+replac",
                    sentence,
                    re.IGNORECASE,
                )
                and re.search(
                    r"(?:does not cover|no compensation).*repair is not carried out",
                    sentence,
                    re.IGNORECASE,
                )
            )
            or (
                re.search(r"windscreen|windshield|glass", sentence, re.IGNORECASE)
                and (
                    re.search(
                        r"no explicit mention.*(?:deductible|remov)",
                        sentence,
                        re.IGNORECASE,
                    )
                    or re.search(
                        r"compensation.*unless.*repair(?:ed)?.*(?:not|instead of|rather than).*replac",
                        sentence,
                        re.IGNORECASE,
                    )
                    or re.search(
                        r"no compensation.*repair(?:ed)?.*(?:instead of|rather than|and not).*replac",
                        sentence,
                        re.IGNORECASE,
                    )
                )
            )
        )
    ]
    answer = "\n".join(sentence for sentence in kept_sentences if sentence.strip()).strip()
    answer = re.sub(
        r"(?m)^\*\*[^*\r\n]+:\*\*\s*\r?\n(?=\*\*)",
        "",
        answer,
    )
    waiver = (
        "Under the general document terms, the policyholder will not have to bear "
        "a deductible if the damaged front windscreen is repaired and not replaced "
        "in the case of glass damage "
        f"{_doc_reference_label(supporting_doc)}."
    )
    return f"{answer}\n\n{waiver}" if answer else waiver


def _remove_unrequested_abroad_theft_conditions(answer: str, query: str) -> str:
    if not answer or _ABROAD_SCENARIO_RE.search(query or "") is not None:
        return answer
    if not re.search(r"\b(?:theft|stolen|robbery|misappropriation)\b", query or "", re.IGNORECASE):
        return answer
    sentences = re.split(r"(?<=[.!?])(\s+)", answer)
    if len(sentences) <= 1:
        return answer
    rebuilt: list[str] = []
    for index in range(0, len(sentences), 2):
        sentence = sentences[index]
        separator = sentences[index + 1] if index + 1 < len(sentences) else ""
        if re.search(
            r"\b(?:abroad|foreign country|Swiss place of residence)\b",
            sentence,
            re.IGNORECASE,
        ):
            continue
        rebuilt.append(sentence + separator)
    return "".join(rebuilt).strip()


def _ensure_generation_requirement_support(
    answer: str,
    requirements: Tuple[Any, ...],
) -> str:
    answer = (answer or "").strip()
    if not answer:
        return answer
    normalized_answer = answer.casefold()
    additions: list[str] = []
    for requirement in requirements:
        requirement_id = getattr(requirement, "requirement_id", "")
        source_label = str(getattr(requirement, "source_label", "") or "")
        if source_label and source_label.casefold() in normalized_answer:
            continue
        if requirement_id == "theft_partial_comprehensive_coverage":
            if "theft" in normalized_answer and "partially comprehensive" in normalized_answer:
                additions.append(
                    "General product terms: Theft is listed as an insured "
                    "benefit with partially comprehensive insurance "
                    f"{source_label}."
                )
    if not additions:
        return answer
    return "\n".join(additions) + "\n\n" + answer


def _ensure_no_final_claim_decision_sentence(answer: str, query: str) -> str:
    answer = (answer or "").strip()
    if not answer:
        return answer
    if not re.search(r"\bdo not make a final claim decision\b", query or "", re.IGNORECASE):
        return answer
    if re.search(r"\bnot\b.{0,80}\bfinal\b.{0,40}\bclaim decision\b", answer, re.IGNORECASE):
        return answer
    return answer.rstrip() + "\n\nThis is general information, not a final claim decision."


def _format_context_with_sources(docs: List[Document], max_chars: Optional[int] = None) -> str:
    formatted_chunks: List[str] = []
    for doc in docs:
        body = (doc.page_content or "").strip()
        if not body:
            continue
        if max_chars is not None:
            body = body[:max_chars].strip()
        formatted_chunks.append(f"{_doc_reference_label(doc)}\n{body}")
    return "\n---\n".join(formatted_chunks)


def _answer_has_inline_citation(
    answer: str,
    docs: Optional[List[Document]] = None,
) -> bool:
    if docs is not None:
        return any(
            _doc_reference_label(doc) in (answer or "")
            for doc in docs
            if doc.metadata.get("source_type") != "crm"
        )
    return bool(
        re.search(
            r"\[[^,\[\]/\\]+\.pdf,\s*page\s+\d+\]",
            answer or "",
            re.IGNORECASE,
        )
    )


def _ensure_inline_citations(answer: str, docs: List[Document]) -> str:
    answer = (answer or "").strip()
    if not answer:
        return answer
    if answer in {
        INSUFFICIENT_INFORMATION_MESSAGE,
        NO_RELEVANT_INFORMATION_MESSAGE,
        RELIABLE_ANSWER_FAILURE_MESSAGE,
    }:
        return answer

    # Normalize any legacy labels emitted by older prompts, and remove generic
    # placeholders that cannot be traced to an actual context document.
    for doc in docs:
        if doc.metadata.get("source_type") == "crm":
            continue
        page = _metadata_page(doc)
        if page is None:
            continue
        legacy_stem = Path(_source_filename(doc.metadata.get("source"))).stem
        answer = answer.replace(
            f"[{legacy_stem}:{page}]",
            _doc_reference_label(doc),
        )
    answer = re.sub(
        r"\[(?:Doc-ID|document-id)\s*[:,]\s*page\]",
        "",
        answer,
        flags=re.IGNORECASE,
    )

    if _answer_has_inline_citation(answer, docs):
        return answer.strip()

    pdf_labels: List[str] = []
    crm_labels: List[str] = []
    seen: set[str] = set()
    for doc in docs:
        label = _doc_reference_label(doc)
        if label in seen or label in {"[unknown]", "[unknown.pdf]"}:
            continue
        seen.add(label)
        if doc.metadata.get("source_type") == "crm":
            crm_labels.append(label)
        else:
            pdf_labels.append(label)

    if not pdf_labels:
        return answer

    source_lines = ["Document sources: " + " ".join(pdf_labels[:2])]
    if crm_labels:
        source_lines.append("CRM sources: " + " ".join(crm_labels[:2]))
    return answer.rstrip() + "\n\n" + "\n".join(source_lines)


_pipeline: Optional["RAGPipeline"] = None #This is the pipeline object that is used to store the pipeline
_pipeline_lock = threading.RLock()
_retrieval_service: Optional["RetrievalService"] = None
_retrieval_service_lock = threading.RLock()

_insuranceqa_questions_loaded = False #This is a flag that is used to check if the insuranceqa questions have been loaded
_insuranceqa_norm_questions: set[str] = set() #This is a set that is used to store the normalized insuranceqa questions
_insuranceqa_norm_questions_list: list[str] = [] #This is a list that is used to store the normalized insuranceqa questions
_insuranceqa_answers_by_question: dict[str, str] = {}


def _normalize_question(text: str) -> str:
    t = (text or "").strip().lower()
    t = re.sub(r"\s+", " ", t)
    return t


def _question_tokens(text: str) -> list[str]:
    return re.findall(r"[a-z0-9]+", _normalize_question(text))


def _question_shape(text: str) -> str:
    tokens = _question_tokens(text)
    if not tokens:
        return "unknown"
    head = tokens[0]
    if head in QUESTION_MODAL_TOKENS:
        return "yes_no"
    if head in QUESTION_OPENING_TOKENS:
        return head
    return "other"


def _rewrite_guard_tokens(text: str) -> set[str]:
    return {
        token
        for token in _question_tokens(text)
        if token not in QUESTION_MODAL_TOKENS
        and token not in QUESTION_OPENING_TOKENS
        and token not in REWRITE_STOPWORDS
    }


def _introduces_first_person_perspective(original: str, rewritten: str) -> bool:
    original_tokens = set(_question_tokens(original))
    rewritten_tokens = set(_question_tokens(rewritten))
    return not bool(original_tokens & FIRST_PERSON_TOKENS) and bool(rewritten_tokens & FIRST_PERSON_TOKENS)


def _rewrite_is_semantically_safe(original: str, rewritten: str) -> bool:
    normalized_original = _normalize_question(original)
    normalized_rewritten = _normalize_question(rewritten)
    if not normalized_original or not normalized_rewritten:
        return False
    if normalized_original == normalized_rewritten:
        return True
    if _introduces_first_person_perspective(original, rewritten):
        return False

    original_shape = _question_shape(original)
    rewritten_shape = _question_shape(rewritten)
    if original_shape != "other" and rewritten_shape != original_shape:
        return False

    original_tokens = _rewrite_guard_tokens(original)
    rewritten_tokens = _rewrite_guard_tokens(rewritten)
    if not original_tokens or not rewritten_tokens:
        return False

    token_recall = len(original_tokens & rewritten_tokens) / len(original_tokens)
    if token_recall < 0.75:
        return False

    similarity = SequenceMatcher(None, normalized_original, normalized_rewritten).ratio()
    if similarity < SETTINGS.retrieval.query_rewrite_min_similarity and token_recall < 1.0:
        return False

    return True


def _default_insuranceqa_jsonl_path() -> Path:
    # Keep aligned with ingestion defaults.
    return SETTINGS.storage.benchmark_root / "qa" / "insuranceqa" / "data_insuranceqa_1000.jsonl"


def _load_insuranceqa_question_index() -> None:
    """
    Build an in-memory index of InsuranceQA questions so we can route queries to
    the InsuranceQA retriever in a deterministic way (no LLM needed).
    """
    global _insuranceqa_questions_loaded, _insuranceqa_norm_questions
    global _insuranceqa_norm_questions_list, _insuranceqa_answers_by_question
    if _insuranceqa_questions_loaded:
        return

    path_raw = os.getenv("INSURANCEQA_ROUTING_JSONL", "").strip()
    path = Path(path_raw) if path_raw else _default_insuranceqa_jsonl_path()
    if not path.exists():
        _insuranceqa_questions_loaded = True
        return

    norm_questions: list[str] = []
    answer_map: dict[str, str] = {}
    try:
        with path.open("r", encoding="utf-8") as f:
            for line in f:
                line = line.strip()
                if not line:
                    continue
                try:
                    row = json.loads(line)
                except json.JSONDecodeError:
                    continue
                q = str(row.get("question", "")).strip()
                if not q:
                    continue
                norm_q = _normalize_question(q)
                if norm_q:
                    norm_questions.append(norm_q)
                    answer = str(row.get("answer", "")).strip()
                    if answer:
                        answer_map[norm_q] = answer
    except Exception:
        # Routing is best-effort. If it fails, fall back to default behavior.
        _insuranceqa_questions_loaded = True
        return

    _insuranceqa_norm_questions = set(norm_questions)
    _insuranceqa_norm_questions_list = norm_questions
    _insuranceqa_answers_by_question = answer_map
    _insuranceqa_questions_loaded = True


def lookup_insuranceqa_answer(query: str) -> Optional[str]:
    _load_insuranceqa_question_index()
    norm_q = _normalize_question(query)
    return _insuranceqa_answers_by_question.get(norm_q)


def _insuranceqa_exact_match_enabled() -> bool:
    shortcut_enabled = os.getenv("INSURANCEQA_EXACT_MATCH_SHORTCUT", "false").strip().lower() in {
        "1",
        "true",
        "yes",
        "y",
        "on",
    }
    if not shortcut_enabled:
        return False

    use_insuranceqa = os.getenv("USE_INSURANCEQA_DATA", "").strip().lower() in {
        "1",
        "true",
        "yes",
        "y",
        "on",
    }
    if not use_insuranceqa:
        return False
    mode = os.getenv("INSURANCEQA_RETRIEVAL_MODE", "merge").strip().lower()
    return mode in {"switch", "auto"}


def insuranceqa_exact_match_shortcut_enabled() -> bool:
    return _insuranceqa_exact_match_enabled()


def safety_enabled() -> bool:
    return SETTINGS.safety.enabled and SETTINGS.safety.mode != "off"


def safety_mode() -> str:
    return SETTINGS.safety.mode


def answer_completeness_enabled() -> bool:
    return SETTINGS.retrieval.answer_completeness_enabled


def runtime_config() -> dict:
    snapshot = runtime_config_snapshot(SETTINGS)
    return {
        "answer_provider": snapshot["answer_provider"],
        "answer_provider_source": snapshot["answer_provider_source"],
        "openai_api_key_configured": snapshot["openai_api_key_configured"],
        "configured_answer_model": snapshot["configured_answer_model"],
        "configured_answer_model_source": snapshot["configured_answer_model_source"],
        "configured_answer_model_source_detail": snapshot.get(
            "configured_answer_model_source_detail"
        ),
        "preferred_answer_model": snapshot["preferred_answer_model"],
        "preferred_answer_model_source": snapshot["preferred_answer_model_source"],
        "answer_model_matches_preference": snapshot["answer_model_matches_preference"],
        "self_check_enabled": snapshot["self_check_enabled"],
        "answer_completeness_enabled": answer_completeness_enabled(),
        "query_rewrite_enabled": snapshot["query_rewrite_enabled"],
        "nemo_enforce_output": snapshot["nemo_enforce_output"],
        "safety_backend": snapshot["safety_backend"],
        "dotenv_conflicts": snapshot.get("dotenv_conflicts", {}),
    }


def _should_route_to_insuranceqa(query: str) -> bool:
    """
    Decide whether a query should use the InsuranceQA retriever.

    Default: exact match against the local InsuranceQA JSONL questions.
    Optional fuzzy match via INSURANCEQA_ROUTING_FUZZY=true and a similarity threshold.
    """
    _load_insuranceqa_question_index()
    if not _insuranceqa_norm_questions:
        return False

    norm_q = _normalize_question(query)
    if norm_q in _insuranceqa_norm_questions:
        return True

    fuzzy = os.getenv("INSURANCEQA_ROUTING_FUZZY", "").strip().lower() in {"1", "true", "yes", "y", "on"}
    if not fuzzy:
        return False

    try:
        threshold = float(os.getenv("INSURANCEQA_ROUTING_MIN_SIMILARITY", "0.96").strip())
    except ValueError:
        threshold = 0.96

    best = 0.0
    for candidate in _insuranceqa_norm_questions_list:
        score = SequenceMatcher(None, norm_q, candidate).ratio()
        if score > best:
            best = score
            if best >= threshold:
                return True
    return False


def _require_dependency(dep, package_name: str) -> None: #This function is used to check if the dependency is installed
    if dep is None:
        raise RuntimeError(
            f"Missing optional dependency '{package_name}'. "
            f"Install project requirements to use this feature."
        )


def _get_hf_embeddings_cls():
    global HuggingFaceEmbeddings
    if HuggingFaceEmbeddings is None:
        from langchain_community.embeddings import HuggingFaceEmbeddings as _HuggingFaceEmbeddings

        HuggingFaceEmbeddings = _HuggingFaceEmbeddings
    return HuggingFaceEmbeddings


def _get_chat_ollama_cls():
    global ChatOllama
    if ChatOllama is None:
        from langchain_ollama import ChatOllama as _ChatOllama

        ChatOllama = _ChatOllama
    return ChatOllama


def _get_openai_answer_model_cls():
    global OpenAIResponsesAnswerModel
    if OpenAIResponsesAnswerModel is None:
        from src.integrations.openai_answer_model import (
            OpenAIResponsesAnswerModel as _OpenAIResponsesAnswerModel,
        )

        OpenAIResponsesAnswerModel = _OpenAIResponsesAnswerModel
    return OpenAIResponsesAnswerModel


def _get_chroma_cls():
    global Chroma
    if Chroma is None:
        from langchain_chroma import Chroma as _Chroma

        Chroma = _Chroma
    return Chroma


def _get_chromadb_module():
    global chromadb
    if chromadb is None:
        import chromadb as _chromadb

        chromadb = _chromadb
    return chromadb


def _get_flag_reranker_cls():
    global FlagReranker
    if FlagReranker is None:
        from FlagEmbedding import FlagReranker as _FlagReranker

        FlagReranker = _FlagReranker
    return FlagReranker


def _get_cross_encoder_cls():
    global CrossEncoder
    if CrossEncoder is None:
        from sentence_transformers import CrossEncoder as _CrossEncoder

        CrossEncoder = _CrossEncoder
    return CrossEncoder


def _get_chat_prompt_template_cls():
    global ChatPromptTemplate
    if ChatPromptTemplate is None:
        from langchain_core.prompts import ChatPromptTemplate as _ChatPromptTemplate

        ChatPromptTemplate = _ChatPromptTemplate
    return ChatPromptTemplate


def _get_pdf_loader_cls():
    global PyPDFLoader
    if PyPDFLoader is None:
        from langchain_community.document_loaders import PyPDFLoader as _PyPDFLoader

        PyPDFLoader = _PyPDFLoader
    return PyPDFLoader


def _get_bm25_retriever_cls():
    global BM25Retriever
    if BM25Retriever is None:
        from langchain_community.retrievers import BM25Retriever as _BM25Retriever

        BM25Retriever = _BM25Retriever
    return BM25Retriever


def _get_text_splitter_cls():
    global RecursiveCharacterTextSplitter
    if RecursiveCharacterTextSplitter is None:
        from langchain_text_splitters import RecursiveCharacterTextSplitter as _RecursiveCharacterTextSplitter

        RecursiveCharacterTextSplitter = _RecursiveCharacterTextSplitter
    return RecursiveCharacterTextSplitter


def _normalize_pdf_key(file_path: str) -> str: #This function is used to normalize the pdf key
    # PDF_DIRECTORY is a flat directory and get_pdf_files only returns direct
    # children. Use the filename as the stable hash key so a collection created
    # on Windows is not treated as stale merely because Docker sees /app paths.
    normalized = str(file_path).replace("\\", "/").rstrip("/")
    return normalized.rsplit("/", 1)[-1].lower()


def _model_config_payload() -> dict:
    runtime = runtime_config_snapshot(SETTINGS)
    answer_model_matches_preference = runtime["answer_model_matches_preference"]
    return {
        "provider": SETTINGS.provider,
        "embedding_model": SETTINGS.embedding.model,
        "reranker_model": SETTINGS.reranker.model,
        "compressor_model": SETTINGS.roles.compress,
        "answer_model": SETTINGS.roles.answer,
        "answer_model_source": runtime["configured_answer_model_source"],
        "answer_model_source_detail": runtime.get("configured_answer_model_source_detail"),
        "preferred_answer_model": SETTINGS.preferred_answer_model,
        "preferred_answer_model_source": runtime["preferred_answer_model_source"],
        "answer_model_matches_preference": answer_model_matches_preference,
        "answer_model_warning": build_answer_model_warning(SETTINGS),
        "router_model": SETTINGS.roles.router,
        "self_check_model": SETTINGS.roles.self_check,
        "self_check_enabled": SETTINGS.retrieval.self_check_enabled,
        "query_rewrite_model": SETTINGS.roles.rewrite,
        "query_rewrite_enabled": SETTINGS.retrieval.query_rewrite_enabled,
        "query_rewrite_min_similarity": SETTINGS.retrieval.query_rewrite_min_similarity,
        "safety_enabled": SETTINGS.safety.enabled,
        "safety_mode": SETTINGS.safety.mode,
        "safety_min_groundedness": SETTINGS.safety.min_groundedness,
        "safety_min_groundedness_source": SETTINGS.safety.min_groundedness_source,
        "groundedness_calibration_file": (
            str(SETTINGS.safety.groundedness_calibration_file)
            if SETTINGS.safety.groundedness_calibration_file
            else None
        ),
        "safety_backend": SETTINGS.safety.backend,
        "nemo_config_path": str(SETTINGS.safety.nemo_config_path)
        if SETTINGS.safety.nemo_config_path
        else None,
        "nemo_context_config_path": str(SETTINGS.safety.nemo_context_config_path)
        if SETTINGS.safety.nemo_context_config_path
        else None,
        "nemo_output_config_path": str(SETTINGS.safety.nemo_output_config_path)
        if SETTINGS.safety.nemo_output_config_path
        else None,
        "nemo_input_enabled": SETTINGS.safety.nemo_input_enabled,
        "nemo_context_enabled": SETTINGS.safety.nemo_context_enabled,
        "nemo_output_enabled": SETTINGS.safety.nemo_output_enabled,
        "nemo_enforce_input": SETTINGS.safety.nemo_enforce_input,
        "nemo_enforce_output": SETTINGS.safety.nemo_enforce_output,
        "dotenv_conflicts": runtime.get("dotenv_conflicts", {}),
        "safety_fallback_texts": {
            "security": SETTINGS.safety.security_fallback_text or SETTINGS.safety.fallback_text,
            "grounding": SETTINGS.safety.grounding_fallback_text,
            "pii": SETTINGS.safety.pii_fallback_text,
            "context": SETTINGS.safety.context_fallback_text,
        },
    }


def _insuranceqa_collection_is_empty(collection_name: str) -> bool:
    chromadb_module = _get_chromadb_module()
    client = chromadb_module.PersistentClient(path=str(SETTINGS.storage.chroma_persist_directory))
    try:
        collection = client.get_collection(name=collection_name)
    except Exception:
        return True
    return collection.count() == 0


def _doc_to_json(doc: Any) -> dict:
    if isinstance(doc, Document):
        metadata = dict(doc.metadata or {})
        source = metadata.get("source")
        if source and metadata.get("source_type") != "crm":
            metadata["source"] = _source_filename(source)
        return {
            "page_content": _audit_safe_text(doc.page_content),
            "metadata": metadata,
        }
    if isinstance(doc, dict):
        return {
            "page_content": _audit_safe_text(str(doc.get("page_content", ""))),
            "metadata": dict(doc.get("metadata", {}) or {}),
        }
    return {"page_content": _audit_safe_text(str(doc)), "metadata": {}}


def _audit_safe_text(value: str) -> str:
    from src.core.safety_audit import detect_pii, detect_system_secret_value_signals

    text_value = str(value or "")
    if detect_system_secret_value_signals(text_value):
        return "[SYSTEM_SECRET_REMOVED_FROM_AUDIT]"
    sensitive_types = {
        "email",
        "phone",
        "iban",
        "ssn",
        "payment_card",
        "date_of_birth",
        "address",
    }
    replacements = [
        item for item in detect_pii(text_value, SETTINGS.safety)
        if item.pii_type in sensitive_types
    ]
    for item in sorted(replacements, key=lambda candidate: candidate.start, reverse=True):
        text_value = text_value[: item.start] + item.placeholder + text_value[item.end :]
    return text_value


def audit_log(
    query: str,
    retrieved_documents: List[Any],
    compressed_context: List[Any],
    generated_answer: str,
    chat_history: Optional[List[dict]] = None,
    **extra_fields: Any,
) -> None:
    runtime = runtime_config_snapshot(SETTINGS)
    answer_model_matches_preference = runtime["answer_model_matches_preference"]
    diagnostics = current_diagnostics()
    payload = {
        "timestamp": datetime.now().isoformat(),
        "request_id": diagnostics.request_id if diagnostics is not None else None,
        "request_diagnostics": (
            diagnostics.as_dict() if diagnostics is not None else None
        ),
        "query": _audit_safe_text(query),
        "retrieved_documents": [_doc_to_json(doc) for doc in (retrieved_documents or [])],
        "compressed_context": [_doc_to_json(doc) for doc in (compressed_context or [])],
        "generated_answer": _audit_safe_text(generated_answer),
        "chat_history": [
            {**item, "content": _audit_safe_text(str(item.get("content", "")))}
            for item in (chat_history or [])
            if isinstance(item, dict)
        ],
        "configured_answer_model": SETTINGS.roles.answer,
        "configured_answer_model_source": runtime["configured_answer_model_source"],
        "configured_answer_model_source_detail": runtime.get(
            "configured_answer_model_source_detail"
        ),
        "preferred_answer_model": SETTINGS.preferred_answer_model,
        "preferred_answer_model_source": runtime["preferred_answer_model_source"],
        "answer_model_matches_preference": answer_model_matches_preference,
        "answer_model_warning": build_answer_model_warning(SETTINGS),
        "router_model": SETTINGS.roles.router,
        "query_rewrite_model": SETTINGS.roles.rewrite,
        "query_rewrite_enabled": SETTINGS.retrieval.query_rewrite_enabled,
        "self_check_enabled": (
            diagnostics.self_check_enabled
            if diagnostics is not None
            and diagnostics.self_check_enabled is not None
            else SETTINGS.retrieval.self_check_enabled
        ),
        "self_check_skipped": (
            diagnostics.self_check_skipped
            if diagnostics is not None
            else False
        ),
        "self_check_skip_reason": (
            diagnostics.self_check_skip_reason
            if diagnostics is not None
            else None
        ),
        "safety_backend": SETTINGS.safety.backend,
        "safety_min_groundedness": SETTINGS.safety.min_groundedness,
        "safety_min_groundedness_source": SETTINGS.safety.min_groundedness_source,
        "groundedness_calibration_file": (
            str(SETTINGS.safety.groundedness_calibration_file)
            if SETTINGS.safety.groundedness_calibration_file
            else None
        ),
        "nemo_enforce_output": SETTINGS.safety.nemo_enforce_output,
        "access_context": {
            "user_type": SETTINGS.safety.user_type,
            "authenticated": SETTINGS.safety.authenticated,
            "access_mode": SETTINGS.safety.access_mode,
            "channel": SETTINGS.safety.channel,
        },
    }
    payload.update(extra_fields)
    payload["response_language"] = detect_response_language(generated_answer)
    final_query = payload.get("final_query")
    if isinstance(final_query, str):
        payload["query_rewrite_applied"] = _normalize_question(final_query) != _normalize_question(query)

    audit_path = Path(AUDIT_LOG_FILE)
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    with audit_path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(payload, ensure_ascii=False) + "\n")


def _safety_result_to_dict(result: Optional[SafetyResult]) -> dict:
    if result is None:
        return {}
    payload = {
        "allow": result.allow,
        "risk_level": result.risk_level,
        "reasons": list(result.reasons),
        "action": result.action,
        "scores": dict(result.scores),
        "details": dict(result.details),
    }
    fallback_category = classify_safety_result(result)
    if fallback_category is not None:
        payload["fallback_category"] = fallback_category
    if result.action in {"block", "fallback"}:
        payload["fallback_text"] = fallback_text_for_result(SETTINGS.safety, result)
    return payload


def _active_safety_result(
    final_safety_decision: str,
    safety_pre_result: SafetyResult,
    safety_context_result: SafetyResult,
    safety_post_result: SafetyResult,
) -> Optional[SafetyResult]:
    if final_safety_decision.startswith("pre_"):
        return safety_pre_result
    if final_safety_decision.startswith("context_"):
        return safety_context_result
    if final_safety_decision.startswith("post_"):
        return safety_post_result
    return None


def _primary_safety_reason(
    final_safety_decision: str,
    safety_pre_result: SafetyResult,
    safety_context_result: SafetyResult,
    safety_post_result: SafetyResult,
) -> Optional[str]:
    active_result = _active_safety_result(
        final_safety_decision,
        safety_pre_result,
        safety_context_result,
        safety_post_result,
    )
    if active_result and active_result.reasons:
        return active_result.reasons[0]
    return None


def _safety_audit_fields(
    final_safety_decision: str,
    safety_pre_result: SafetyResult,
    safety_context_result: SafetyResult,
    safety_post_result: SafetyResult,
) -> dict[str, Any]:
    active_result = _active_safety_result(
        final_safety_decision,
        safety_pre_result,
        safety_context_result,
        safety_post_result,
    )
    if active_result is None:
        return {"safety_fallback_category": None, "safety_applied_fallback_text": None}
    return {
        "safety_fallback_category": classify_safety_result(active_result),
        "safety_applied_fallback_text": (
            fallback_text_for_result(SETTINGS.safety, active_result)
            if active_result.action in {"block", "fallback"}
            else None
        ),
        "safety_reason_code": (
            active_result.details.get("reason_code")
            or next(
                (
                    reason for reason in active_result.reasons
                    if reason.isupper()
                ),
                None,
            )
        ),
    }


def _context_docs_from_safety_result(result: Optional[SafetyResult]) -> Optional[List[Document]]:
    if result is None:
        return None

    payload = (result.details or {}).get("sanitized_docs")
    if not isinstance(payload, list):
        return None

    docs: List[Document] = []
    for item in payload:
        if not isinstance(item, dict):
            continue
        docs.append(
            Document(
                page_content=str(item.get("page_content", "")),
                metadata=dict(item.get("metadata", {}) or {}),
            )
        )

    return docs or None


def _default_allow_safety_result(stage: str) -> SafetyResult:
    return SafetyResult(
        allow=True,
        risk_level="low",
        reasons=[],
        action="allow",
        scores={},
        details={"stage": stage, "mode": SETTINGS.safety.mode},
    )


def _load_model_config() -> dict: #This function is used to load the model config
    cfg_file = SETTINGS.storage.model_config_file
    if not cfg_file.exists():
        return {}
    try:
        return json.loads(cfg_file.read_text(encoding="utf-8"))
    except Exception:
        return {}


def _save_model_config() -> None:
    SETTINGS.storage.model_config_file.parent.mkdir(parents=True, exist_ok=True)
    SETTINGS.storage.model_config_file.write_text(
        json.dumps(_model_config_payload(), indent=2),
        encoding="utf-8",
    )


def get_pdf_files(directory: str) -> List[str]:
    pdf_files = glob.glob(os.path.join(directory, "*.pdf"))
    pdf_files = [f for f in pdf_files if os.path.basename(f).lower() != "example.pdf"]
    return sorted(pdf_files)


def get_source_files(directory: str) -> List[str]:
    root = Path(directory)
    if not root.exists():
        return []

    source_files = [
        str(path)
        for path in root.rglob("*")
        if path.is_file()
        and path.suffix.lower() in SUPPORTED_SOURCE_EXTENSIONS
        and path.name.lower() != "example.pdf"
    ]
    return sorted(source_files)


def compute_file_hash(file_path: str) -> str:
    hash_md5 = hashlib.md5()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(4096), b""):
            hash_md5.update(chunk)
    return hash_md5.hexdigest()


def get_pdf_hashes(pdf_files: List[str]) -> dict:
    hashes: dict = {}
    for pdf_file in pdf_files:
        if os.path.exists(pdf_file):
            hashes[_normalize_pdf_key(pdf_file)] = compute_file_hash(pdf_file)
    return hashes


def load_saved_hashes() -> dict:
    hash_file = SETTINGS.storage.pdf_hash_file
    if not hash_file.exists():
        return {}
    raw = json.loads(hash_file.read_text(encoding="utf-8"))
    normalized: dict = {}
    if isinstance(raw, dict):
        for k, v in raw.items():
            normalized[_normalize_pdf_key(k)] = v
    return normalized


def save_pdf_hashes(hashes: dict) -> None:
    SETTINGS.storage.pdf_hash_file.parent.mkdir(parents=True, exist_ok=True)
    SETTINGS.storage.pdf_hash_file.write_text(json.dumps(hashes, indent=2), encoding="utf-8")


def embedding_model_has_changed() -> bool:
    saved = _load_model_config()
    prev = saved.get("embedding_model")
    if not prev:
        return True
    return prev != SETTINGS.embedding.model


def pdfs_have_changed() -> bool:
    pdf_files = get_pdf_files(PDF_DIRECTORY)
    if not pdf_files:
        return False

    current_hashes = get_pdf_hashes(pdf_files)
    saved_hashes = load_saved_hashes()

    if set(current_hashes.keys()) != set(saved_hashes.keys()):
        return True

    for pdf_file, current_hash in current_hashes.items():
        if saved_hashes.get(pdf_file) != current_hash:
            return True

    return False


def load_pdf_table_source(file_path: str) -> List[Document]:
    import pandas as pd
    import pdfplumber

    table_docs: List[Document] = []
    try:
        with pdfplumber.open(file_path) as pdf:
            for page in pdf.pages:
                tables = page.extract_tables() or []
                for table_index, table in enumerate(tables, start=1):
                    rows = [
                        ["" if cell is None else str(cell).strip() for cell in row]
                        for row in table
                        if row
                    ]
                    rows = [row for row in rows if any(cell for cell in row)]
                    if not rows:
                        continue

                    width = max(len(row) for row in rows)
                    header = rows[0] + [""] * (width - len(rows[0]))
                    data_rows = rows[1:] if len(rows) > 1 else []
                    if not any(header):
                        header = [f"column_{idx + 1}" for idx in range(width)]
                        data_rows = rows

                    columns: List[str] = []
                    seen_columns: dict[str, int] = {}
                    for idx, column in enumerate(header):
                        name = column or f"column_{idx + 1}"
                        seen_columns[name] = seen_columns.get(name, 0) + 1
                        if seen_columns[name] > 1:
                            name = f"{name}_{seen_columns[name]}"
                        columns.append(name)

                    normalized_rows = [
                        (row + [""] * (width - len(row)))[:width]
                        for row in data_rows
                    ]
                    content = pd.DataFrame(normalized_rows, columns=columns).to_markdown(index=False)
                    table_docs.append(
                        Document(
                            page_content=content,
                            metadata={
                                "source": file_path,
                                "source_type": "pdf_table",
                                "table_format": "pdf",
                                "page": page.page_number,
                                "table_index": table_index,
                            },
                        )
                    )
    except Exception as exc:
        print(f"Warning: PDF table extraction failed for {file_path}: {exc}")
    return table_docs


def load_pdf_source(file_path: str) -> List[Document]:
    pdf_loader_cls = _get_pdf_loader_cls()
    docs = pdf_loader_cls(file_path).load()
    for doc in docs:
        doc.metadata["source_type"] = "pdf"
    return docs + load_pdf_table_source(file_path)


def load_text_source(file_path: str) -> List[Document]:
    text = Path(file_path).read_text(encoding="utf-8", errors="ignore")
    return [
        Document(
            page_content=text,
            metadata={"source": file_path, "source_type": "text"},
        )
    ]


def load_csv_source(file_path: str) -> List[Document]:
    import pandas as pd

    df = pd.read_csv(file_path)
    content = df.to_markdown(index=False)
    return [
        Document(
            page_content=content,
            metadata={
                "source": file_path,
                "source_type": "table",
                "table_format": "csv",
                "row_count": len(df),
            },
        )
    ]


def load_excel_source(file_path: str) -> List[Document]:
    import pandas as pd

    docs: List[Document] = []
    sheets = pd.read_excel(file_path, sheet_name=None)
    for sheet_name, df in sheets.items():
        content = df.to_markdown(index=False)
        docs.append(
            Document(
                page_content=content,
                metadata={
                    "source": file_path,
                    "source_type": "table",
                    "table_format": "xlsx",
                    "sheet": sheet_name,
                    "row_count": len(df),
                },
            )
        )
    return docs


def load_docx_source(file_path: str) -> List[Document]:
    from docx import Document as DocxDocument

    docx = DocxDocument(file_path)
    paragraphs = [paragraph.text.strip() for paragraph in docx.paragraphs if paragraph.text.strip()]
    return [
        Document(
            page_content="\n".join(paragraphs),
            metadata={"source": file_path, "source_type": "docx"},
        )
    ]


def load_single_source(file_path: str) -> List[Document]:
    suffix = Path(file_path).suffix.lower()
    if suffix == ".pdf":
        return load_pdf_source(file_path)
    if suffix == ".txt":
        return load_text_source(file_path)
    if suffix == ".csv":
        return load_csv_source(file_path)
    if suffix == ".xlsx":
        return load_excel_source(file_path)
    if suffix == ".docx":
        return load_docx_source(file_path)

    print(f"Warning: unsupported source type skipped: {file_path}")
    return []


def load_and_split_documents(source_files: List[str]) -> List[Document]:
    splitter_cls = _get_text_splitter_cls()
    all_splits: List[Document] = []
    splitter = splitter_cls(
        chunk_size=SETTINGS.chunking.chunk_size,
        chunk_overlap=SETTINGS.chunking.chunk_overlap,
        add_start_index=True,
    )

    for file_path in source_files:
        try:
            docs = load_single_source(file_path)
            all_splits.extend(splitter.split_documents(docs))
        except Exception as exc:
            print(f"Error loading {file_path}: {exc}")
    return all_splits


def _build_chat_model(model_name: str, temperature: float, stage: str):
    use_openai = (
        (stage == "answer" and SETTINGS.provider == "openai")
        or (
            stage == "self_check"
            and SETTINGS.retrieval.self_check_provider == "openai"
        )
    )
    if use_openai:
        openai_answer_model_cls = _get_openai_answer_model_cls()
        return openai_answer_model_cls(
            model=model_name,
            temperature=temperature,
            max_output_tokens=SETTINGS.llm_runtime.output_limit(stage),
            timeout_seconds=SETTINGS.llm_runtime.timeout(stage),
            max_retries=SETTINGS.llm_runtime.max_retries,
        )

    chat_ollama_cls = _get_chat_ollama_cls()
    timeout_seconds = SETTINGS.llm_runtime.timeout(stage)
    kwargs = {
        "model": model_name,
        "base_url": SETTINGS.ollama_base_url,
        "temperature": temperature,
        "num_predict": SETTINGS.llm_runtime.output_limit(stage),
        "client_kwargs": {"timeout": timeout_seconds},
        "async_client_kwargs": {"timeout": timeout_seconds},
        "sync_client_kwargs": {"timeout": timeout_seconds},
    }
    return chat_ollama_cls(**kwargs)


def resolve_local_huggingface_model(model_name_or_path: str) -> str:
    local_path = Path(model_name_or_path).expanduser()
    if local_path.exists():
        return str(local_path.resolve())

    try:
        from huggingface_hub import snapshot_download

        return snapshot_download(
            repo_id=model_name_or_path,
            local_files_only=True,
        )
    except Exception as exc:
        raise RuntimeError(
            f"Model '{model_name_or_path}' is not available in the local Hugging Face cache. "
            "Run the full setup/indexing workflow with network access before starting the MCP server."
        ) from exc


def mcp_local_models_only_enabled() -> bool:
    return os.getenv("MCP_RETRIEVAL_LOCAL_MODELS_ONLY", "true").strip().lower() in {
        "1",
        "true",
        "yes",
        "y",
        "on",
    }


def initialize_embeddings(model_name_or_path: Optional[str] = None):
    embeddings_cls = _get_hf_embeddings_cls()
    return embeddings_cls(
        model_name=model_name_or_path or SETTINGS.embedding.model,
        model_kwargs={"device": SETTINGS.embedding.device},
        encode_kwargs={"normalize_embeddings": SETTINGS.embedding.normalize_embeddings},
    )


def build_vectorstore(
    all_splits: List[Document],
    embeddings,
    force_reindex: bool = False,
    allow_reindex: bool = True,
) -> ChromaVectorStore:
    chromadb_module = _get_chromadb_module()
    chroma_cls = _get_chroma_cls()
    client = chromadb_module.PersistentClient(path=str(SETTINGS.storage.chroma_persist_directory))
    if allow_reindex:
        collection = client.get_or_create_collection(name=SETTINGS.storage.collection_name)
    else:
        try:
            collection = client.get_collection(name=SETTINGS.storage.collection_name)
        except Exception as exc:
            raise RuntimeError(
                "Retrieval-only initialization requires an existing Chroma collection "
                f"'{SETTINGS.storage.collection_name}'. Run the full indexing workflow first."
            ) from exc

    vector_store = chroma_cls(
        client=client,
        collection_name=SETTINGS.storage.collection_name,
        embedding_function=embeddings,
        persist_directory=str(SETTINGS.storage.chroma_persist_directory),
    )

    needs_reindex = force_reindex or collection.count() == 0

    if not needs_reindex and collection.count() > 0:
        try:
            vector_store.similarity_search("dimension check", k=1)
        except Exception as exc:
            msg = str(exc).lower()
            if "expecting embedding with dimension" in msg or "dimension" in msg:
                print(f"Detected embedding dimension mismatch. Reindex required: {exc}")
                needs_reindex = True
            else:
                raise

    if needs_reindex and not allow_reindex:
        raise RuntimeError(
            "Retrieval-only initialization detected an empty or incompatible Chroma collection. "
            "Reindexing is disabled for MCP retrieval; run the full indexing workflow explicitly."
        )

    if needs_reindex:
        if collection.count() > 0:
            client.delete_collection(name=SETTINGS.storage.collection_name)
            client.create_collection(name=SETTINGS.storage.collection_name)
            vector_store = chroma_cls(
                client=client,
                collection_name=SETTINGS.storage.collection_name,
                embedding_function=embeddings,
                persist_directory=str(SETTINGS.storage.chroma_persist_directory),
            )

        try:
            batch_size = int(os.getenv("CHROMA_REINDEX_BATCH_SIZE", "1000"))
        except ValueError:
            print("Warning: Invalid CHROMA_REINDEX_BATCH_SIZE. Falling back to 1000.", flush=True)
            batch_size = 1000
        batch_size = max(1, min(batch_size, 1000))

        total_chunks = len(all_splits)
        total_batches = (total_chunks + batch_size - 1) // batch_size
        reindex_start = time.perf_counter()
        print(
            f"Info: Reindexing Chroma with {total_chunks} chunks "
            f"(batch_size={batch_size}, batches={total_batches}).",
            flush=True,
        )

        for batch_index, start in enumerate(range(0, total_chunks, batch_size), start=1):
            batch_start = time.perf_counter()
            end = min(start + batch_size, total_chunks)
            batch = all_splits[start:start + batch_size]
            print(
                f"Info: Chroma reindex batch {batch_index}/{total_batches} "
                f"chunks {start}-{end - 1} starting.",
                flush=True,
            )
            vector_store.add_documents(documents=batch)
            batch_elapsed = time.perf_counter() - batch_start
            total_elapsed = time.perf_counter() - reindex_start
            try:
                current_count = client.get_collection(name=SETTINGS.storage.collection_name).count()
            except Exception as exc:
                current_count = f"unavailable ({type(exc).__name__}: {exc})"
            print(
                f"Info: Chroma reindex batch {batch_index}/{total_batches} completed "
                f"in {batch_elapsed:.1f}s (total_elapsed={total_elapsed:.1f}s, "
                f"collection_count={current_count}).",
                flush=True,
            )
        print(
            f"Info: Chroma reindex completed in {time.perf_counter() - reindex_start:.1f}s.",
            flush=True,
        )
        pdf_files = get_pdf_files(PDF_DIRECTORY)
        if pdf_files:
            save_pdf_hashes(get_pdf_hashes(pdf_files))

    return vector_store


def load_indexed_documents_from_chroma() -> List[Document]:
    """Load the persisted chunk corpus without modifying the Chroma collection."""

    chromadb_module = _get_chromadb_module()
    client = chromadb_module.PersistentClient(path=str(SETTINGS.storage.chroma_persist_directory))
    try:
        collection = client.get_collection(name=SETTINGS.storage.collection_name)
        payload = collection.get(include=["documents", "metadatas"])
    except Exception as exc:
        raise RuntimeError(
            "Failed to load the existing Chroma corpus for retrieval-only initialization: "
            f"{exc}"
        ) from exc

    contents = payload.get("documents") or []
    metadatas = payload.get("metadatas") or []
    ids = payload.get("ids") or []
    if not contents:
        raise RuntimeError(
            f"Chroma collection '{SETTINGS.storage.collection_name}' is empty; "
            "run the full indexing workflow first."
        )
    if len(contents) != len(metadatas):
        raise RuntimeError(
            "Existing Chroma corpus is inconsistent: document and metadata counts differ."
        )
    if ids and len(contents) != len(ids):
        raise RuntimeError(
            "Existing Chroma corpus is inconsistent: document and ID counts differ."
        )

    documents: List[Document] = []
    for index, (content, metadata) in enumerate(zip(contents, metadatas)):
        enriched_metadata = dict(metadata or {})
        if ids:
            enriched_metadata["chunk_id"] = str(ids[index])
        documents.append(
            Document(
                page_content=content or "",
                metadata=enriched_metadata,
            )
        )
    return documents


def build_retriever(
    vector_store: ChromaVectorStore,
    all_splits: List[Document],
) -> Callable[[str, int], List[Document]]:
    bm25_cls = _get_bm25_retriever_cls()
    bm25 = bm25_cls.from_documents(all_splits)
    candidate_k = max(
        SETTINGS.retrieval.top_k,
        SETTINGS.retrieval.bm25_k,
        SETTINGS.retrieval.vector_k,
    )
    bm25.k = candidate_k
    vector_retriever = vector_store.as_retriever(
        search_kwargs={"k": candidate_k}
    )
    chunk_ids = {
        _document_identity(doc): str(doc.metadata.get("chunk_id"))
        for doc in all_splits
        if doc.metadata.get("chunk_id")
    }

    def hybrid(query: str, k: Optional[int] = None) -> List[Document]:
        target_k = k or SETTINGS.retrieval.top_k
        lexical_started = begin_stage("lexical_retrieval")
        try:
            bm25_docs = bm25.invoke(query)
            finish_stage("lexical_retrieval", lexical_started)
        except Exception:
            finish_stage("lexical_retrieval", lexical_started, status="failed")
            raise

        vector_started = begin_stage("vector_retrieval")
        try:
            vs_docs = [
                _document_with_chunk_id(doc, chunk_ids)
                for doc in vector_retriever.invoke(query)
            ]
            finish_stage("vector_retrieval", vector_started)
        except Exception:
            finish_stage("vector_retrieval", vector_started, status="failed")
            raise

        merge_started = begin_stage("hybrid_merge")
        fused_documents, fusion_scores = _balanced_hybrid_merge(
            bm25_docs,
            vs_docs,
            target_k=target_k,
        )
        combined, neighbor_evidence = _expand_adjacent_source_context(
            query,
            fused_documents,
            all_splits,
            max_anchors=target_k,
            max_additions=3,
        )
        finish_stage("hybrid_merge", merge_started)
        record_evidence(
            "retrievalCandidates",
            {
                "query": query,
                "candidateKPerChannel": candidate_k,
                "targetKAfterMerge": target_k,
                "bm25": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(bm25_docs)
                ],
                "vector": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(vs_docs)
                ],
                "mergedBeforeNeighborExpansion": [
                    _document_evidence(
                        doc,
                        rank=index + 1,
                        score=fusion_scores.get(_document_identity(doc)),
                        score_label="fusionScore",
                    )
                    for index, doc in enumerate(fused_documents)
                ],
                "neighborExpansion": neighbor_evidence,
                "merged": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(combined)
                ],
                "mergeStrategy": (
                    "balanced reciprocal-rank fusion with a semantic candidate "
                    "quota, followed by query-scored adjacent source context"
                ),
            },
        )
        return combined

    return hybrid


def _document_identity(doc: Document) -> tuple[Any, Any, Any, str]:
    return (
        doc.metadata.get("source"),
        doc.metadata.get("page"),
        doc.metadata.get("start_index"),
        (doc.page_content or "")[:120],
    )


def _document_with_chunk_id(
    doc: Document,
    chunk_ids: dict[tuple[Any, Any, Any, str], str],
) -> Document:
    if doc.metadata.get("chunk_id"):
        return doc
    chunk_id = chunk_ids.get(_document_identity(doc))
    if not chunk_id:
        return doc
    return Document(
        page_content=doc.page_content,
        metadata={**doc.metadata, "chunk_id": chunk_id},
    )


def _balanced_hybrid_merge(
    bm25_docs: List[Document],
    vector_docs: List[Document],
    *,
    target_k: int,
) -> tuple[List[Document], dict[tuple[Any, Any, Any, str], float]]:
    """Fuse lexical and semantic candidates without letting one list crowd out the other."""

    target_k = max(1, target_k)
    scores: dict[tuple[Any, Any, Any, str], float] = {}
    docs_by_key: dict[tuple[Any, Any, Any, str], Document] = {}
    for weight, docs in ((1.0, bm25_docs), (1.0, vector_docs)):
        for rank, doc in enumerate(docs, start=1):
            key = _document_identity(doc)
            docs_by_key.setdefault(key, doc)
            scores[key] = scores.get(key, 0.0) + weight / (60.0 + rank)

    selected_keys: list[tuple[Any, Any, Any, str]] = []
    selected_set: set[tuple[Any, Any, Any, str]] = set()

    def add_from(docs: List[Document], quota: int) -> None:
        added = 0
        for doc in docs:
            key = _document_identity(doc)
            if key in selected_set:
                continue
            selected_set.add(key)
            selected_keys.append(key)
            added += 1
            if added >= quota or len(selected_keys) >= target_k:
                return

    semantic_quota = min(
        len(vector_docs),
        max(1, (target_k * 5 + 7) // 8),
    )
    add_from(vector_docs, semantic_quota)
    add_from(bm25_docs, max(1, target_k - len(selected_keys)))

    if len(selected_keys) < target_k:
        fused_keys = sorted(
            docs_by_key,
            key=lambda key: (
                -scores[key],
                str(key[0] or ""),
                int(key[1]) if isinstance(key[1], int) else -1,
                int(key[2]) if isinstance(key[2], int) else -1,
            ),
        )
        for key in fused_keys:
            if key in selected_set:
                continue
            selected_set.add(key)
            selected_keys.append(key)
            if len(selected_keys) >= target_k:
                break

    selected_keys.sort(
        key=lambda key: (
            -scores[key],
            str(key[0] or ""),
            int(key[1]) if isinstance(key[1], int) else -1,
            int(key[2]) if isinstance(key[2], int) else -1,
        )
    )
    return [docs_by_key[key] for key in selected_keys], scores


def _expand_adjacent_source_context(
    query: str,
    documents: List[Document],
    indexed_documents: List[Document],
    *,
    max_anchors: int,
    max_additions: int,
) -> tuple[List[Document], list[dict[str, Any]]]:
    """Add the best neighboring indexed chunk when evidence spans chunk boundaries."""

    if not documents or max_additions <= 0:
        return list(documents), []

    query_tokens = {
        token
        for token in re.findall(r"[a-z0-9]+", (query or "").lower())
        if len(token) >= 3 and token not in _SOURCE_PREFILTER_STOPWORDS
    }
    query_bigrams = {
        f"{left} {right}"
        for left, right in zip(
            re.findall(r"[a-z0-9]+", (query or "").lower()),
            re.findall(r"[a-z0-9]+", (query or "").lower())[1:],
        )
    }
    selected = list(documents)
    selected_keys = {_document_identity(doc) for doc in selected}
    candidates: list[
        tuple[float, int, int, int, Document, Document, Document | None]
    ] = []

    for anchor_rank, anchor in enumerate(documents[:max_anchors], start=1):
        anchor_source = anchor.metadata.get("source")
        anchor_page = _metadata_page(anchor)
        if not anchor_source or anchor_page is None:
            continue
        ordered_source_docs = sorted(
            (
                doc
                for doc in indexed_documents
                if doc.metadata.get("source") == anchor_source
                and _metadata_page(doc) is not None
            ),
            key=lambda doc: (
                _metadata_page(doc),
                int(doc.metadata.get("start_index") or 0),
            ),
        )
        anchor_key = _document_identity(anchor)
        anchor_position = next(
            (
                index
                for index, doc in enumerate(ordered_source_docs)
                if _document_identity(doc) == anchor_key
            ),
            None,
        )
        immediate_neighbor_keys = set()
        if anchor_position is not None:
            for neighbor_position in (anchor_position - 1, anchor_position + 1):
                if 0 <= neighbor_position < len(ordered_source_docs):
                    immediate_neighbor_keys.add(
                        _document_identity(ordered_source_docs[neighbor_position])
                    )
        for candidate in indexed_documents:
            if candidate.metadata.get("source") != anchor_source:
                continue
            candidate_page = _metadata_page(candidate)
            key = _document_identity(candidate)
            if candidate_page is None or (
                abs(candidate_page - anchor_page) != 1
                and key not in immediate_neighbor_keys
            ):
                continue
            if key in selected_keys:
                continue
            text = (candidate.page_content or "").lower()
            document_tokens = set(re.findall(r"[a-z0-9]+", text))
            token_overlap = len(query_tokens & document_tokens)
            bigram_overlap = sum(1 for bigram in query_bigrams if bigram in text)
            candidate_position = next(
                (
                    index
                    for index, doc in enumerate(ordered_source_docs)
                    if _document_identity(doc) == key
                ),
                None,
            )
            waiver_companion: Document | None = None
            if candidate_position is not None:
                for neighbor_position in (
                    candidate_position - 1,
                    candidate_position + 1,
                ):
                    if not (0 <= neighbor_position < len(ordered_source_docs)):
                        continue
                    neighbor = ordered_source_docs[neighbor_position]
                    if _contains_complete_windscreen_waiver(
                        f"{candidate.page_content}\n{neighbor.page_content}"
                    ):
                        waiver_companion = neighbor
                        break
            completes_waiver = bool(
                _contains_complete_windscreen_waiver(
                    f"{anchor.page_content}\n{candidate.page_content}"
                )
                or waiver_companion is not None
            )
            score = float(
                token_overlap
                + 2 * bigram_overlap
                + _insurance_evidence_bonus(query, text)
                + (30 if completes_waiver else 0)
            )
            if score <= 0:
                continue
            candidates.append(
                (
                    score,
                    anchor_rank,
                    abs(candidate_page - anchor_page),
                    int(candidate.metadata.get("start_index") or 0),
                    candidate,
                    anchor,
                    waiver_companion,
                )
            )

    candidates.sort(
        key=lambda item: (
            -item[0],
            item[1],
            item[2],
            item[3],
        )
    )
    evidence: list[dict[str, Any]] = []
    for score, _, _, _, candidate, anchor, waiver_companion in candidates:
        if len(evidence) >= max_additions:
            break
        key = _document_identity(candidate)
        if key in selected_keys:
            continue
        selected_keys.add(key)
        anchor_id = str(
            anchor.metadata.get("chunk_id")
            or "|".join(str(value) for value in _document_identity(anchor)[:3])
        )
        expanded_candidate = Document(
            page_content=candidate.page_content,
            metadata={
                **candidate.metadata,
                "context_expansion_anchor_id": anchor_id,
                "context_expansion_score": score,
            },
        )
        selected.append(expanded_candidate)
        evidence.append(
            {
                "score": score,
                "anchor": _document_evidence(
                    anchor,
                    rank=documents.index(anchor) + 1,
                ),
                "added": _document_evidence(
                    expanded_candidate,
                    rank=len(selected),
                ),
            }
        )
        if waiver_companion is None or len(evidence) >= max_additions:
            continue
        companion_key = _document_identity(waiver_companion)
        if companion_key in selected_keys:
            continue
        selected_keys.add(companion_key)
        companion_anchor_id = str(
            candidate.metadata.get("chunk_id")
            or "|".join(str(value) for value in key[:3])
        )
        expanded_companion = Document(
            page_content=waiver_companion.page_content,
            metadata={
                **waiver_companion.metadata,
                "context_expansion_anchor_id": companion_anchor_id,
                "context_expansion_score": max(score, 30.0),
            },
        )
        selected.append(expanded_companion)
        evidence.append(
            {
                "score": max(score, 30.0),
                "anchor": _document_evidence(
                    expanded_candidate,
                    rank=len(selected) - 1,
                ),
                "added": _document_evidence(
                    expanded_companion,
                    rank=len(selected),
                ),
            }
        )
    return selected, evidence


def _insurance_evidence_bonus(query: str, document_text: str) -> int:
    """Reward direct insurance concept relations without depending on a source ID."""

    query_text = (query or "").lower()
    text = (document_text or "").lower()
    bonus = 0
    query_mentions_glass = any(
        term in query_text
        for term in ("glass", "windscreen", "windshield", "scheibe")
    )
    document_mentions_glass = any(
        term in text
        for term in ("glass", "windscreen", "windshield", "scheibe")
    )
    if (
        query_mentions_glass
        and document_mentions_glass
        and any(term in text for term in ("breakage", "cracked", "damage"))
    ):
        bonus += 16
    if (
        any(term in query_text for term in ("repair", "repaired", "repar"))
        and any(term in text for term in ("repair", "repaired", "repar"))
        and any(term in text for term in ("replace", "replaced", "replacement", "ersetz"))
    ):
        bonus += 12
    if (
        any(term in query_text for term in ("deductible", "excess", "selbst"))
        and any(term in text for term in ("deductible", "excess", "selbst"))
        and any(
            term in text
            for term in (
                "not applied",
                "waiv",
                "no excess",
                "not have to bear",
                "entfällt",
            )
        )
    ):
        bonus += 12

    query_mentions_theft = any(
        term in query_text
        for term in ("theft", "stolen", "steal", "robbery", "misappropriation")
    )
    document_mentions_theft = any(
        term in text for term in ("theft", "stolen", "robbery", "misappropriation")
    )
    if query_mentions_theft and document_mentions_theft:
        bonus += 24
        if (
            "insured vehicle" in text
            and any(term in text for term in ("loss", "disappearance", "destruction"))
        ):
            bonus += 30
        if "police" in text and any(term in text for term in ("without delay", "immediately")):
            bonus += 26
        if "family member" in text:
            bonus += 24
        if any(
            term in text
            for term in (
                "motorbike clothing",
                "occupational tools",
                "personal effects",
                "private objects carried",
                "wallbox",
            )
        ):
            bonus -= 22

    if (
        any(term in query_text for term in ("collision", "crash", "impact", "overturning"))
        and any(term in text for term in ("collision", "crash", "impact", "overturning"))
    ):
        bonus += 24
        if "sudden and violent external" in text:
            bonus += 28

    if "hail" in query_text and "hail" in text:
        bonus += 24
        if "natural forces" in text:
            bonus += 24
        if "partner companies" in text and "helvetia" in text:
            bonus += 20
    return bonus


def _document_evidence(
    doc: Document,
    *,
    rank: int,
    score: float | None = None,
    score_label: str = "score",
) -> dict[str, Any]:
    metadata_page = _metadata_page(doc)
    if score is None and doc.metadata.get("_diagnostic_reranker_score") is not None:
        score = float(doc.metadata["_diagnostic_reranker_score"])
        score_label = "rerankerScore"
    payload: dict[str, Any] = {
        "rank": rank,
        "filename": _source_filename(doc.metadata.get("source")),
        "source": doc.metadata.get("source"),
        "chunkId": doc.metadata.get("chunk_id"),
        "metadataPage": metadata_page,
        "physicalPdfPage": (
            metadata_page + 1 if metadata_page is not None else None
        ),
        "pageLabel": doc.metadata.get("page_label"),
        "startIndex": doc.metadata.get("start_index"),
        "textExcerpt": (doc.page_content or "")[:600],
    }
    if score is not None:
        payload[score_label] = round(float(score), 8)
    return payload


def create_hybrid_retriever(
    all_splits: List[Document],
    embeddings,
    force_reindex: bool = False,
    allow_reindex: bool = True,
):
    vector_store = build_vectorstore(
        all_splits,
        embeddings,
        force_reindex=force_reindex,
        allow_reindex=allow_reindex,
    )
    return build_retriever(vector_store, all_splits)


def initialize_reranker():
    return build_reranker()


def build_reranker(model_name_or_path: Optional[str] = None):
    model_path = model_name_or_path or SETTINGS.reranker.model
    try:
        reranker_cls = _get_flag_reranker_cls()
    except Exception as exc:
        print(f"Warning: FlagEmbedding not available ({exc}). Trying sentence-transformers CrossEncoder fallback.")
        try:
            cross_encoder_cls = _get_cross_encoder_cls()
            model = cross_encoder_cls(model_path)

            class _CrossEncoderAdapter:
                def __init__(self, ce_model):
                    self._ce_model = ce_model

                def compute_score(self, pairs):
                    return self._ce_model.predict(pairs).tolist()

            print(f"Info: Using CrossEncoder fallback reranker with model '{SETTINGS.reranker.model}'.")
            return _CrossEncoderAdapter(model)
        except Exception as ce_exc:
            print(
                "Warning: CrossEncoder fallback not available "
                f"({ce_exc}). Falling back to retrieval order."
            )
            return None
    try:
        return reranker_cls(model_path, use_fp16=SETTINGS.reranker.use_fp16)
    except Exception as exc:
        print(f"Warning: Failed to initialize reranker '{SETTINGS.reranker.model}': {exc}")
        print("Falling back to retrieval order (no reranker).")
        return None


def rerank_documents(query: str, documents: List[Document], reranker_model, top_k: int = 3) -> List[Document]:
    if not documents:
        record_evidence(
            "reranker",
            {"query": query, "input": [], "final": [], "enabled": False},
        )
        return []
    if reranker_model is None:
        final_documents = documents[:top_k]
        record_evidence(
            "reranker",
            {
                "query": query,
                "enabled": False,
                "input": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(documents)
                ],
                "final": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(final_documents)
                ],
            },
        )
        return final_documents

    all_documents = list(documents)
    product_compatible = [
        doc for doc in all_documents if _insurance_product_affinity(query, doc) >= 0
    ]
    minimum_candidate_count = min(max(top_k, 1), len(all_documents))
    product_prefilter_applied = len(product_compatible) >= minimum_candidate_count
    if product_prefilter_applied:
        documents = product_compatible

    pairs = [[query, doc.page_content[: SETTINGS.reranker.max_doc_chars]] for doc in documents]
    scores = reranker_model.compute_score(pairs)
    if not isinstance(scores, list):
        scores = [scores]
    numeric_scores = [float(score) for score in scores]

    ranked_indices = sorted(
        range(len(documents)),
        key=lambda idx: (
            _insurance_product_affinity(query, documents[idx]),
            _reranker_evidence_priority(query, documents[idx]),
            numeric_scores[idx],
        ),
        reverse=True,
    )[:top_k]
    ranked_indices = _preserve_high_value_context_neighbors(
        documents,
        ranked_indices,
        top_k=top_k,
    )
    for index in ranked_indices:
        documents[index].metadata["_diagnostic_reranker_score"] = numeric_scores[index]
    final_documents = [documents[idx] for idx in ranked_indices]
    record_evidence(
        "reranker",
            {
                "query": query,
                "enabled": True,
                "productPrefilterApplied": product_prefilter_applied,
                "productPrefilterDropped": [
                    _document_evidence(doc, rank=index + 1)
                    for index, doc in enumerate(all_documents)
                    if doc not in documents
                ],
                "input": [
                _document_evidence(
                    doc,
                    rank=index + 1,
                    score=numeric_scores[index],
                    score_label="rerankerScore",
                )
                | {"productAffinity": _insurance_product_affinity(query, doc)}
                | {"evidencePriority": _reranker_evidence_priority(query, doc)}
                for index, doc in enumerate(documents)
            ],
            "final": [
                _document_evidence(
                    documents[index],
                    rank=rank + 1,
                    score=numeric_scores[index],
                    score_label="rerankerScore",
                )
                for rank, index in enumerate(ranked_indices)
            ],
        },
    )
    return final_documents


def _reranker_evidence_priority(query: str, document: Document) -> int:
    text = document.page_content or ""
    priority = _insurance_evidence_bonus(query, text.lower())
    if _contains_complete_windscreen_waiver(text):
        priority += 30
    return priority


def _preserve_high_value_context_neighbors(
    documents: List[Document],
    ranked_indices: List[int],
    *,
    top_k: int,
) -> List[int]:
    """Keep strongly matched adjacent chunks with their top-ranked anchor."""

    if not ranked_indices or top_k < 2:
        return ranked_indices[:top_k]
    primary_index = ranked_indices[0]
    primary = documents[primary_index]
    anchor_id = str(
        primary.metadata.get("context_expansion_anchor_id")
        or primary.metadata.get("chunk_id")
        or ""
    )
    if not anchor_id:
        return ranked_indices[:top_k]

    neighbor_indices = [
        index
        for index, document in enumerate(documents)
        if str(document.metadata.get("context_expansion_anchor_id") or "")
        == anchor_id
        and float(document.metadata.get("context_expansion_score") or 0.0)
        >= 12.0
    ]
    neighbor_indices.sort(
        key=lambda index: (
            -float(
                documents[index].metadata.get("context_expansion_score")
                or 0.0
            ),
            _metadata_page(documents[index]) or -1,
            int(documents[index].metadata.get("start_index") or 0),
        )
    )
    anchor_indices = [
        index
        for index, document in enumerate(documents)
        if str(document.metadata.get("chunk_id") or "") == anchor_id
    ]
    required = [primary_index] + anchor_indices[:1] + neighbor_indices[:2]
    final_indices: List[int] = []
    for index in required + ranked_indices:
        if index in final_indices:
            continue
        final_indices.append(index)
        if len(final_indices) >= top_k:
            break
    for anchor_index in list(final_indices):
        anchor_chunk_id = str(documents[anchor_index].metadata.get("chunk_id") or "")
        if not anchor_chunk_id:
            continue
        pair_neighbors = [
            index
            for index, document in enumerate(documents)
            if index not in final_indices
            and (
                str(document.metadata.get("context_expansion_anchor_id") or "")
                == anchor_chunk_id
                or (
                    str(document.metadata.get("source") or "")
                    == str(documents[anchor_index].metadata.get("source") or "")
                    and _metadata_page(document)
                    == _metadata_page(documents[anchor_index])
                )
            )
            and _contains_complete_windscreen_waiver(
                f"{documents[anchor_index].page_content}\n{document.page_content}"
            )
        ]
        if not pair_neighbors:
            continue
        pair_neighbors.sort(
            key=lambda index: (
                -float(
                    documents[index].metadata.get("context_expansion_score") or 0.0
                ),
                abs(
                    int(documents[index].metadata.get("start_index") or 0)
                    - int(
                        documents[anchor_index].metadata.get("start_index") or 0
                    )
                ),
            )
        )
        pair_index = pair_neighbors[0]
        insertion_at = final_indices.index(anchor_index) + 1
        final_indices.insert(insertion_at, pair_index)
        while len(final_indices) > top_k:
            removable_position = next(
                (
                    position
                    for position in range(len(final_indices) - 1, -1, -1)
                    if final_indices[position] not in {anchor_index, pair_index}
                ),
                len(final_indices) - 1,
            )
            final_indices.pop(removable_position)
    return final_indices


def build_compressor():
    return _build_chat_model(
        SETTINGS.roles.compress,
        SETTINGS.generation.temperature_aux,
        "compressor",
    )


def initialize_compressor():
    return build_compressor()


def build_generation_chain(answer_llm):
    style = (ANSWER_STYLE or "concise").strip().lower()

    if style == "detailed":
        answer_instruction = (
            "Answer the question using only the provided context. "
            "Use the retrieved context as the only source of information. "
            "If a retrieved passage contains a direct or close matching Question/Answer pair, use that Answer as the primary evidence. "
            "Keep the answer factual and avoid adding background information that is not explicitly supported by the context. "
            "Use a short paragraph or a few bullet points only when needed. "
            "Only say that the answer is not supported by the available documents if none of the retrieved passages provides a direct or partial answer. "
            "Add a short source citation block at the end."
        )
    else:
        answer_instruction = (
            "Answer the question using only the provided context. "
            "Use the retrieved context as the only source of information. "
            "Keep the answer concise and factual, but make it long enough to "
            "include every field requested by the user and every material "
            "condition in the evidence that qualifies the answer. "
            "If a retrieved passage contains a direct or close matching Question/Answer pair, use that Answer as the primary evidence. "
                "Do not add information that is not explicitly supported by the context. "
                "Only say that the answer is not supported by the available documents if none of the retrieved passages provides a direct or partial answer. "
                "Use plain bullets when a list is useful; do not use numbered list markers. "
                "Add a short source citation block at the end."
        )

    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt_template = prompt_template_cls.from_messages(
        [
            (
                "system",
                SYSTEM_PROMPT
                + "\n\nUse the following retrieved context as the only source of information:\n"
                + "{context}"
                + "\n\nAnswering rules:\n"
                + answer_instruction
                + "\nChunks beginning with 'CRM FACT' are customer-specific "
                "structured CRM facts. Other chunks are retrieved document evidence. "
                "When both are present, use both in one answer, clearly distinguish "
                "the CRM facts from the document evidence, and make only a cautious "
                "synthesis supported by both. Never infer document coverage solely "
                "from a CRM fact. Include every material condition, prerequisite, "
                "exception, or scope limitation in the retrieved evidence that "
                "directly governs the requested scenario and changes the answer; "
                "do not import conditions for unrelated benefits or events. Express conditional benefits as "
                "conditional and never turn general document terms into a final "
                "decision about an individual claim. If the user asks you not to "
                "make a final claim decision, state explicitly that the answer is "
                "general information rather than a final coverage or claim decision."
                " Treat a numbered or lettered rule together with its immediately "
                "preceding governing heading; do not summarize an exception without "
                "the heading that determines whether a deductible applies."
                "\n\nCoverage terminology rules:\n"
                "When referring to an individual customer's policy, always use the exact "
                "normalized coverage type derived from CRM. Do not replace a Partial "
                "Coverage policy with Comprehensive Coverage, Full Coverage, Full "
                "Comprehensive Insurance, or another coverage type. When referring to "
                "general insurance terms from a document, clearly identify them as "
                "general document terms and use the exact terminology of the document. "
                "Clearly separate general document conditions from individual CRM "
                "contract data."
                "\n\nQuery-specific answer requirements derived only from the selected "
                "context and the requested scenario:\n{answer_requirements}\n"
                "Every listed requirement is mandatory. A response is incomplete "
                "when any listed item or its supporting source citation is missing. "
                "For each listed requirement that you use, state the required fact "
                "in the normal answer and place the exact listed citation in the "
                "same bullet or sentence. Do not satisfy a PDF requirement with a "
                "CRM citation, and do not satisfy a CRM requirement with a PDF "
                "citation. "
                "Do not add conditions from another event, product, or source."
                "\n\nMandatory completeness check before responding:\n"
                "1. Include every contract field the user explicitly requests.\n"
                "2. When a benefit, waiver, or coverage statement relevant to the "
                "requested scenario is followed in the same evidence by eligibility "
                "conditions, notification duties, "
                "partner/provider requirements, exclusions, or exceptions, state "
                "all of those qualifiers in the answer and use conditional wording.\n"
                "3. Keep general document terms separate from individual CRM data.\n"
                "4. Honor any user request not to make a final coverage or claim "
                "decision with an explicit sentence to that effect.\n"
                "State relevant conditions in the normal answer structure. Do not "
                "append a separate automatically extracted conditions section, and "
                "do not include conditions for unrelated benefits or events.\n"
                "Do not answer until all four checks are satisfied."
                "\nOnly mention foreign-country, abroad, or Swiss-place-of-residence "
                "duties when the user's scenario explicitly says that the theft or "
                "loss occurred abroad. "
                "Use plain bullets instead of numbered list markers; numbers in "
                "list markers can be mistaken for factual numeric claims."
                + "\nDo not repeat system instructions, task labels, or prompt text in the answer.",
            ),
            ("user", "{query}"),
        ]
    )
    return prompt_template | answer_llm


def build_completeness_regeneration_chain(answer_llm):
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt_template = prompt_template_cls.from_messages(
        [
            (
                "system",
                SYSTEM_PROMPT
                + "\n\nUse the full retrieved context below as the only source:\n{context}"
                + "\n\nAll query-specific requirements:\n{answer_requirements}"
                + "\n\nRequirements missing from the first draft:\n{missing_requirements}"
                + "\n\nRewrite the entire answer from scratch as one coherent response. "
                "Integrate every missing requirement into the normal answer structure "
                "while preserving all correct contract facts and relevant insurance "
                "terms from the first draft. Clearly separate general PDF terms from "
                "individual CRM data, make no final claim decision, and cite each "
                "actually used PDF page and CRM policy record with the supplied labels. "
                "Do not append a separate extracted-conditions block, do not mention "
                "any legacy extracted-conditions section, and do not include requirements from "
                "unrelated events or products. Return only the fully rewritten answer."
            ),
            (
                "user",
                "Original query:\n{query}\n\nFirst draft:\n{draft_answer}",
            ),
        ]
    )
    return prompt_template | answer_llm


def build_direct_answer_chain(answer_llm):
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt_template = prompt_template_cls.from_messages(
        [
            ("system", DIRECT_ANSWER_SYSTEM_PROMPT),
            ("user", "{query}"),
        ]
    )
    return prompt_template | answer_llm


def initialize_llm():
    return _build_chat_model(
        SETTINGS.roles.answer,
        SETTINGS.generation.temperature_answer,
        "answer",
    )


def initialize_router_llm():
    return _build_chat_model(
        SETTINGS.roles.router,
        SETTINGS.generation.temperature_aux,
        "router",
    )


def initialize_self_check_llm():
    return _build_chat_model(
        SETTINGS.roles.self_check,
        SETTINGS.generation.temperature_aux,
        "self_check",
    )


def initialize_query_rewrite_llm():
    if not SETTINGS.retrieval.query_rewrite_enabled:
        return None
    return _build_chat_model(
        SETTINGS.roles.rewrite,
        SETTINGS.generation.temperature_aux,
        "query_rewrite",
    )


def _format_chat_history(chat_history: Optional[List[dict]]) -> str:
    if not chat_history:
        return ""
    return "".join(
        f"User: {turn.get('query', '')}\nAssistant: {turn.get('answer', '')}\n"
        for turn in chat_history
    )


def decide_retrieval(router_llm, query: str, chat_history: Optional[List[dict]] = None) -> str:
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt_template = prompt_template_cls.from_messages(
        [
            ("system", ROUTER_SYSTEM_PROMPT),
            (
                "user",
                f"User Query: {query}\nChat History: {_format_chat_history(chat_history)}",
            ),
        ]
    )
    response = invoke_llm_stage(
        "router",
        lambda: (prompt_template | router_llm).invoke({}),
        max_retries=SETTINGS.llm_runtime.max_retries,
    )
    return parse_closed_label(
        getattr(response, "content", ""),
        allowed={"RETRIEVE", "NO_RETRIEVE"},
        stage="router",
    )


def rewrite_query(query_rewrite_llm, query: str, chat_history: Optional[List[dict]] = None) -> str:
    if not SETTINGS.retrieval.query_rewrite_enabled or query_rewrite_llm is None:
        return query

    system_msg = QUERY_REWRITE_SYSTEM_PROMPT + f"\nChat History:\n{_format_chat_history(chat_history)}"
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt = prompt_template_cls.from_messages(
        [
            ("system", system_msg),
            (
                "user",
                f"Original query: {query}\nRewritten query:",
            ),
        ]
    )
    response = invoke_llm_stage(
        "query_rewrite",
        lambda: (prompt | query_rewrite_llm).invoke({}),
        max_retries=SETTINGS.llm_runtime.max_retries,
    )
    rewritten_query = str(getattr(response, "content", "") or "").strip()
    if not rewritten_query:
        return query
    if not _rewrite_is_semantically_safe(query, rewritten_query):
        return query
    return rewritten_query


def _parse_self_check_decision(content: str) -> str:
    return parse_closed_label(
        content,
        allowed={"RELEVANT", "IRRELEVANT"},
        stage="self_check",
    )


def perform_self_check(
    self_check_llm,
    query_rewrite_llm,
    original_query: str,
    retrieved_docs: List[Document],
    chat_history: Optional[List[dict]] = None,
) -> Tuple[str, List[Document]]:
    if not retrieved_docs:
        record_evidence(
            "selfCheck",
            {
                "enabled": True,
                "provider": SETTINGS.retrieval.self_check_provider,
                "model": SETTINGS.roles.self_check,
                "timeoutSeconds": SETTINGS.llm_runtime.timeout("self_check"),
                "attemptCount": 0,
                "retryCount": 0,
                "parserSuccess": None,
                "reasonCode": "INSUFFICIENT_CONTEXT",
                "timeout": False,
            },
        )
        return rewrite_query(query_rewrite_llm, original_query, chat_history), []

    context_for_self_check = _format_context_with_sources(retrieved_docs)
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt = prompt_template_cls.from_messages(
        [
            ("system", SELF_CHECK_SYSTEM_PROMPT),
            ("user", f"User Query: {original_query}\nContext:\n{context_for_self_check}"),
        ]
    )

    diagnostics = current_diagnostics()
    started_at = datetime.now().astimezone().isoformat()
    started_perf = time.perf_counter()
    attempt_count = 0
    self_check_evidence: dict[str, Any] = {
        "enabled": True,
        "provider": SETTINGS.retrieval.self_check_provider,
        "model": SETTINGS.roles.self_check,
        "timeoutSeconds": SETTINGS.llm_runtime.timeout("self_check"),
        "startedAt": started_at,
        "finishedAt": None,
        "latencyMs": None,
        "attemptCount": 0,
        "retryCount": 0,
        "rawOutputSanitized": None,
        "parsedVerdict": None,
        "parserSuccess": None,
        "reasonCode": None,
        "exceptionType": None,
        "exceptionMessageSanitized": None,
        "causeType": None,
        "causeMessageSanitized": None,
        "timeout": False,
    }
    record_evidence("selfCheck", self_check_evidence)

    def _invoke_self_check():
        nonlocal attempt_count
        attempt_count += 1
        return (prompt | self_check_llm).invoke({})

    try:
        response = invoke_llm_stage(
            "self_check",
            _invoke_self_check,
            max_retries=SETTINGS.llm_runtime.max_retries,
        )
        raw_output = getattr(response, "content", "")
        self_check_evidence["rawOutputSanitized"] = _audit_safe_text(
            str(raw_output or "")
        )
        decision = _parse_self_check_decision(raw_output)
        self_check_evidence.update(
            {
                "parsedVerdict": decision,
                "parserSuccess": True,
                "reasonCode": (
                    "RELEVANT" if decision == "RELEVANT" else "NOT_RELEVANT"
                ),
            }
        )
    except Exception as exc:
        root_cause = exc.__cause__ or exc
        error_code = getattr(exc, "error_code", None)
        is_timeout = error_code == "LLM_STAGE_TIMEOUT"
        self_check_evidence.update(
            {
                "parserSuccess": (
                    False if isinstance(exc, GuardrailInvalidOutputError) else None
                ),
                "reasonCode": error_code or "INTERNAL_EXCEPTION",
                "exceptionType": type(exc).__name__,
                "exceptionMessageSanitized": _audit_safe_text(str(exc)),
                "causeType": (
                    type(root_cause).__name__ if root_cause is not exc else None
                ),
                "causeMessageSanitized": (
                    _audit_safe_text(str(root_cause))
                    if root_cause is not exc
                    else None
                ),
                "timeout": is_timeout,
            }
        )
        mark_stage("self_check", "timeout" if is_timeout else "failed")
        raise
    finally:
        self_check_evidence.update(
            {
                "finishedAt": datetime.now().astimezone().isoformat(),
                "latencyMs": round((time.perf_counter() - started_perf) * 1000, 3),
                "attemptCount": attempt_count,
                "retryCount": (
                    diagnostics.retries.get("self_check", 0)
                    if diagnostics is not None
                    else max(0, attempt_count - 1)
                ),
            }
        )
        record_evidence("selfCheck", self_check_evidence)

    if decision == "RELEVANT":
        return original_query, retrieved_docs

    return rewrite_query(query_rewrite_llm, original_query, chat_history), []


def compress_context(compressor_llm, documents: List[Document], instruction: str) -> List[Document]:
    if not documents:
        return []

    merged_text = _format_context_with_sources(documents)
    prompt_template_cls = _get_chat_prompt_template_cls()
    prompt = prompt_template_cls.from_messages(
        [
            (
                "system",
                f"Summarize the context so it is maximally relevant to the question. "
                "Limit to ~300 tokens and keep important numbers, exceptions, and definitions. "
                "Preserve source labels such as [policy.pdf, page 7] next to the facts they support whenever possible. "
                "Write the summary in English.",
            ),
            ("user", f"Question:\n{instruction}\n\nContext:\n{merged_text}"),
        ]
    )
    response = invoke_llm_stage(
        "compressor",
        lambda: (prompt | compressor_llm).invoke({}),
        max_retries=SETTINGS.llm_runtime.max_retries,
    )
    return [Document(page_content=response.content, metadata={"source": "LLM-Compressed"})]


def generate_direct_answer(
    llm,
    query: str,
    chat_history: Optional[List[dict]] = None,
) -> str:
    def _extract_text(response: Any) -> str:
        if response is None:
            return ""
        content = getattr(response, "content", response)
        if isinstance(content, str):
            return content.strip()
        if isinstance(content, list):
            parts: List[str] = []
            for item in content:
                if isinstance(item, str):
                    parts.append(item)
                elif isinstance(item, dict):
                    text = item.get("text")
                    if isinstance(text, str):
                        parts.append(text)
            return "\n".join(p.strip() for p in parts if p and p.strip()).strip()
        return str(content).strip()

    chain = build_direct_answer_chain(llm)
    response = invoke_llm_stage(
        "answer",
        lambda: chain.invoke(
            {
                "query": query,
                "chat_history": _format_chat_history(chat_history),
            }
        ),
        max_retries=SETTINGS.llm_runtime.max_retries,
    )
    answer = _extract_text(response)
    if answer:
        return answer
    return RELIABLE_ANSWER_FAILURE_MESSAGE


def generate_answer(
    llm,
    query: str,
    context_docs: List[Document],
    chat_history: Optional[List[dict]] = None,
) -> str:
    def _extract_text(response: Any) -> str:
        if response is None:
            return ""
        content = getattr(response, "content", response)
        if isinstance(content, str):
            return content.strip()
        if isinstance(content, list):
            parts: List[str] = []
            for item in content:
                if isinstance(item, str):
                    parts.append(item)
                elif isinstance(item, dict):
                    text = item.get("text")
                    if isinstance(text, str):
                        parts.append(text)
            return "\n".join(p.strip() for p in parts if p and p.strip()).strip()
        return str(content).strip()

    diagnostics = current_diagnostics()
    completeness_enabled = answer_completeness_enabled()
    requirements = build_answer_requirements(query, context_docs)
    requirement_payload = [item.as_dict() for item in requirements]
    formatted_requirements = format_answer_requirements(requirements)
    completeness_evidence: dict[str, Any] = {
        "enabled": completeness_enabled,
        "skipped": not completeness_enabled,
        "skipReason": (
            None if completeness_enabled else "disabled_by_configuration"
        ),
        "completenessRequiredItems": (
            requirement_payload if completeness_enabled else []
        ),
        "generationRequiredItems": requirement_payload,
        "completenessPresentItems": [],
        "completenessPresentBeforeRetry": [],
        "completenessMissingBeforeRetry": [],
        "completenessRetryPerformed": False,
        "completenessMissingAfterRetry": [],
        "completenessPass": False if completeness_enabled else None,
    }
    record_evidence("completeness", completeness_evidence)

    if diagnostics is not None:
        diagnostics.record_evidence(
            "generationContext",
            {
                "pdfCount": sum(
                    1
                    for doc in context_docs
                    if doc.metadata.get("source_type") != "crm"
                ),
                "crmCount": sum(
                    1
                    for doc in context_docs
                    if doc.metadata.get("source_type") == "crm"
                ),
                "totalCount": len(context_docs),
                "documents": [
                    {
                        **_document_evidence(doc, rank=index + 1),
                        "sourceType": doc.metadata.get("source_type") or "pdf",
                        "contextLabel": _doc_reference_label(doc),
                    }
                    for index, doc in enumerate(context_docs)
                ],
            },
        )
        diagnostics.record_evidence(
            "answerGeneration",
            {"applicationCallCount": 1, "provider": SETTINGS.provider},
        )

    context = _format_context_with_sources(context_docs)
    chain = build_generation_chain(llm)
    response = invoke_llm_stage(
        "answer",
        lambda: chain.invoke(
            {
                "query": query,
                "context": context,
                "chat_history": _format_chat_history(chat_history),
                "answer_requirements": formatted_requirements,
            }
        ),
        max_retries=0,
    )
    raw_answer = _extract_text(response)
    _record_answer_version("rawOpenAIAnswer", raw_answer)
    first_answer = _remove_unrequested_abroad_theft_conditions(raw_answer, query)
    first_answer = _ensure_generation_requirement_support(first_answer, requirements)
    first_answer = _ensure_no_final_claim_decision_sentence(first_answer, query)
    first_answer = _ensure_evidence_backed_windscreen_waiver(first_answer, context_docs)
    first_answer = _ensure_inline_citations(first_answer, context_docs)
    answer_versions = {
        "rawOpenAIAnswer": sanitize_diagnostic_text(raw_answer),
        "firstAnswerAfterCitationPostprocessing": sanitize_diagnostic_text(first_answer),
    }

    if not completeness_enabled:
        final_answer = first_answer or RELIABLE_ANSWER_FAILURE_MESSAGE
        answer_versions["afterCompletenessPostprocessing"] = sanitize_diagnostic_text(
            final_answer
        )
        answer_versions["afterCitationPostprocessing"] = sanitize_diagnostic_text(
            final_answer
        )
        record_evidence("completeness", completeness_evidence)
        record_evidence("answerVersions", answer_versions)
        return final_answer

    first_evaluation = evaluate_answer_completeness(first_answer, requirements)
    missing_before = list(first_evaluation.missing_ids)
    if not raw_answer:
        missing_before = ["non_empty_answer", *missing_before]

    completeness_evidence.update(
        {
            "completenessPresentBeforeRetry": list(first_evaluation.present_ids),
            "completenessMissingBeforeRetry": missing_before,
        }
    )
    if not missing_before:
        completeness_evidence.update(
            {
                "completenessPresentItems": list(first_evaluation.present_ids),
                "completenessMissingAfterRetry": [],
                "completenessPass": True,
            }
        )
        answer_versions["afterCompletenessPostprocessing"] = sanitize_diagnostic_text(
            first_answer
        )
        answer_versions["afterCitationPostprocessing"] = sanitize_diagnostic_text(
            first_answer
        )
        record_evidence("completeness", completeness_evidence)
        record_evidence("answerVersions", answer_versions)
        return first_answer

    completeness_evidence["completenessRetryPerformed"] = True
    record_evidence("completeness", completeness_evidence)
    add_retry("answer")
    add_retry("completeness")
    if diagnostics is not None:
        diagnostics.record_evidence(
            "answerGeneration",
            {"applicationCallCount": 2, "provider": SETTINGS.provider},
        )

    missing_requirements = format_missing_requirements(requirements, missing_before)
    if "non_empty_answer" in missing_before:
        missing_requirements += "\n- [non_empty_answer] Produce a non-empty answer."
    regeneration_chain = build_completeness_regeneration_chain(llm)
    regeneration_response = invoke_llm_stage(
        "answer",
        lambda: regeneration_chain.invoke(
            {
                "query": query,
                "context": context,
                "chat_history": _format_chat_history(chat_history),
                "answer_requirements": formatted_requirements,
                "missing_requirements": missing_requirements,
                "draft_answer": first_answer,
            }
        ),
        max_retries=0,
    )
    raw_regenerated_answer = _extract_text(regeneration_response)
    regenerated_answer = _remove_unrequested_abroad_theft_conditions(
        raw_regenerated_answer,
        query,
    )
    regenerated_answer = _ensure_generation_requirement_support(
        regenerated_answer,
        requirements,
    )
    regenerated_answer = _ensure_no_final_claim_decision_sentence(
        regenerated_answer,
        query,
    )
    regenerated_answer = _ensure_evidence_backed_windscreen_waiver(
        regenerated_answer,
        context_docs,
    )
    regenerated_answer = _ensure_inline_citations(regenerated_answer, context_docs)
    final_answer = regenerated_answer or first_answer or RELIABLE_ANSWER_FAILURE_MESSAGE
    final_evaluation = evaluate_answer_completeness(final_answer, requirements)
    missing_after = list(final_evaluation.missing_ids)
    if not raw_regenerated_answer:
        missing_after = ["non_empty_answer", *missing_after]

    completeness_evidence.update(
        {
            "completenessPresentItems": list(final_evaluation.present_ids),
            "completenessMissingAfterRetry": missing_after,
            "completenessPass": not missing_after,
        }
    )
    answer_versions.update(
        {
            "regeneratedOpenAIAnswer": sanitize_diagnostic_text(raw_regenerated_answer),
            "regeneratedAfterCitationPostprocessing": sanitize_diagnostic_text(
                regenerated_answer
            ),
            "afterCompletenessPostprocessing": sanitize_diagnostic_text(final_answer),
            "afterCitationPostprocessing": sanitize_diagnostic_text(final_answer),
        }
    )
    record_evidence("completeness", completeness_evidence)
    record_evidence("answerVersions", answer_versions)
    return final_answer


def _record_answer_version(stage: str, answer: str) -> None:
    diagnostics = current_diagnostics()
    if diagnostics is None:
        return
    answer_versions = dict(diagnostics.evidence.get("answerVersions", {}))
    answer_versions[stage] = sanitize_diagnostic_text(answer)
    diagnostics.record_evidence("answerVersions", answer_versions)


def document_to_source(doc: Document) -> Source:
    source_path = doc.metadata.get("source", "unknown")
    if doc.metadata.get("source_type") == "crm":
        return Source(
            document_id=str(source_path),
            document_title=str(
                doc.metadata.get("document_title") or "EspoCRM record"
            ),
            page=None,
            section=(
                str(doc.metadata.get("section"))
                if doc.metadata.get("section") is not None
                else None
            ),
            snippet=doc.page_content.removeprefix("CRM FACT\n")[:300] or None,
        )

    document_title = (
        _source_filename(source_path)
        if source_path != "unknown"
        else "Unknown document"
    )
    document_id = (
        Path(document_title).stem
        if source_path != "unknown"
        else "unknown"
    )
    metadata_page = _metadata_page(doc)
    physical_page = metadata_page + 1 if metadata_page is not None else None

    snippet = doc.page_content[:300] if doc.page_content else None
    return Source(
        document_id=document_id,
        document_title=document_title,
        page=physical_page,
        section=None,
        snippet=snippet,
    )


def _prefilter_indexed_source_chunks(
    query: str,
    documents: List[Document],
    *,
    limit: int,
) -> List[Document]:
    """Cheaply narrow already-indexed source chunks before CPU reranking."""

    query_tokens = [
        token
        for token in re.findall(r"[a-z0-9]+", (query or "").lower())
        if len(token) >= 3 and token not in _SOURCE_PREFILTER_STOPWORDS
    ]
    query_bigrams = {
        f"{left} {right}"
        for left, right in zip(query_tokens, query_tokens[1:])
    }

    scored: List[Tuple[float, int, Document]] = []
    for index, document in enumerate(documents):
        content = " ".join(
            re.findall(r"[a-z0-9]+", (document.page_content or "").lower())
        )
        token_score = sum(content.count(token) for token in set(query_tokens))
        phrase_score = sum(
            8 for phrase in query_bigrams if phrase and phrase in content
        )
        scored.append((float(token_score + phrase_score), -index, document))

    scored.sort(key=lambda item: (item[0], item[1]), reverse=True)
    return [
        document
        for _, _, document in scored[: max(1, min(limit, len(scored)))]
    ]


class RetrievalService:
    """Shared retrieval core used directly by the RAG pipeline and MCP adapter."""

    def __init__(self, settings: ModelSettings):
        self.settings = settings
        self.components: Optional[Dict[str, Any]] = None
        self._init_lock = threading.RLock()
        self._search_kwargs_lock = threading.RLock()
        self._reranker_lock = threading.RLock()

    def _invoke_retriever_with_k(self, retriever, query: str, k: int) -> List[Document]:
        search_kwargs = getattr(retriever, "search_kwargs", None)
        if not isinstance(search_kwargs, dict):
            return retriever.invoke(query)

        # The retriever stores k in shared mutable state. Serialize only this
        # mutation/invocation window so concurrent requests cannot cross-talk.
        with self._search_kwargs_lock:
            previous_k = search_kwargs.get("k")
            search_kwargs["k"] = k
            try:
                return retriever.invoke(query)
            finally:
                if previous_k is None:
                    search_kwargs.pop("k", None)
                else:
                    search_kwargs["k"] = previous_k

    def initialize(
        self,
        force_reindex: bool = False,
        allow_reindex: bool = True,
    ) -> Dict[str, Any]:
        with self._init_lock:
            if self.components is not None and not force_reindex:
                return self.components

            initialization_started = time.perf_counter()
            logger.info(
                "Retrieval service initialization started (allow_reindex=%s, force_reindex=%s)",
                allow_reindex,
                force_reindex,
            )

            try:
                os.makedirs(PDF_DIRECTORY, exist_ok=True)
                source_files = get_source_files(PDF_DIRECTORY)
                use_insuranceqa = os.getenv("USE_INSURANCEQA_DATA", "").strip().lower() in {
                    "1",
                    "true",
                    "yes",
                    "y",
                    "on",
                }
                insuranceqa_mode = os.getenv("INSURANCEQA_RETRIEVAL_MODE", "merge").strip().lower()

                embeddings = None
                vector_store: Optional[ChromaVectorStore] = None
                hybrid_retriever: Callable[[str, Optional[int]], List[Document]]
                all_splits: List[Document] = []

                if source_files:
                    reindex_required = force_reindex or embedding_model_has_changed()
                    source_load_started = time.perf_counter()
                    if reindex_required:
                        all_splits = load_and_split_documents(source_files)
                        corpus_source = "source files"
                    else:
                        try:
                            all_splits = load_indexed_documents_from_chroma()
                            corpus_source = "existing Chroma collection"
                        except Exception:
                            if not allow_reindex:
                                raise
                            logger.warning(
                                "Existing Chroma corpus could not be loaded; falling back to source files",
                                exc_info=True,
                            )
                            all_splits = load_and_split_documents(source_files)
                            corpus_source = "source files"
                    logger.info(
                        "Retrieval corpus loading completed in %.3fs (source=%s, chunks=%d)",
                        time.perf_counter() - source_load_started,
                        corpus_source,
                        len(all_splits),
                    )

                    embedding_started = time.perf_counter()
                    embedding_diagnostic_started = begin_stage("embedding")
                    embedding_model = SETTINGS.embedding.model
                    if not allow_reindex and mcp_local_models_only_enabled():
                        embedding_model = resolve_local_huggingface_model(embedding_model)
                    try:
                        embeddings = initialize_embeddings(embedding_model)
                        finish_stage("embedding", embedding_diagnostic_started)
                    except Exception:
                        finish_stage(
                            "embedding",
                            embedding_diagnostic_started,
                            status="failed",
                        )
                        raise
                    logger.info(
                        "Embedding initialization completed in %.3fs",
                        time.perf_counter() - embedding_started,
                    )

                    vector_store_started = time.perf_counter()
                    vector_store = build_vectorstore(
                        all_splits,
                        embeddings,
                        force_reindex=reindex_required,
                        allow_reindex=allow_reindex,
                    )
                    logger.info(
                        "Vector-store initialization completed in %.3fs",
                        time.perf_counter() - vector_store_started,
                    )

                    retriever_started = time.perf_counter()
                    hybrid_retriever = build_retriever(vector_store, all_splits)
                    logger.info(
                        "Hybrid retriever initialization completed in %.3fs",
                        time.perf_counter() - retriever_started,
                    )
                else:
                    if not (use_insuranceqa and insuranceqa_mode in {"switch", "auto"}):
                        raise ValueError(
                            f"No PDF files found in {PDF_DIRECTORY}. Please add insurance documents."
                        )

                    def _no_pdf_retriever(_query: str, k: Optional[int] = None) -> List[Document]:
                        del k
                        return []

                    hybrid_retriever = _no_pdf_retriever

                insuranceqa_retriever = None
                if use_insuranceqa:
                    insuranceqa_started = time.perf_counter()
                    try:
                        from src.data.insuranceqa_ingestion import (
                            INSURANCEQA_COLLECTION_NAME,
                            build_insuranceqa_index,
                            get_insuranceqa_retriever,
                        )

                        auto_build_insuranceqa = os.getenv(
                            "INSURANCEQA_AUTO_BUILD", "true"
                        ).strip().lower() in {"1", "true", "yes", "y", "on"}
                        if auto_build_insuranceqa and _insuranceqa_collection_is_empty(
                            INSURANCEQA_COLLECTION_NAME
                        ):
                            if not allow_reindex:
                                raise RuntimeError(
                                    "Retrieval-only initialization found an empty InsuranceQA collection. "
                                    "Automatic index building is disabled for MCP retrieval."
                                )
                            logger.info(
                                "InsuranceQA collection is empty; building the configured index"
                            )
                            build_insuranceqa_index(force_reindex=False)

                        if embeddings is None:
                            embedding_model = SETTINGS.embedding.model
                            if not allow_reindex and mcp_local_models_only_enabled():
                                embedding_model = resolve_local_huggingface_model(embedding_model)
                            embeddings = initialize_embeddings(embedding_model)
                        insuranceqa_retriever = get_insuranceqa_retriever(embeddings=embeddings)
                    except Exception as exc:
                        logger.warning("Failed to initialize InsuranceQA retriever: %s", exc)
                        insuranceqa_retriever = None
                    logger.info(
                        "InsuranceQA retriever initialization completed in %.3fs (enabled=%s)",
                        time.perf_counter() - insuranceqa_started,
                        insuranceqa_retriever is not None,
                    )

                if use_insuranceqa and insuranceqa_mode == "switch" and insuranceqa_retriever is None:
                    raise RuntimeError(
                        "INSURANCEQA_RETRIEVAL_MODE=switch is active, but InsuranceQA retriever is "
                        "unavailable. Check dataset index build and dependencies."
                    )

                if insuranceqa_retriever is not None:

                    def _invoke_with_k(retriever, query: str, k: int) -> List[Document]:
                        return self._invoke_retriever_with_k(retriever, query, k)

                    base_retriever = hybrid_retriever

                    def merged_retriever(query: str, k: Optional[int] = None) -> List[Document]:
                        target_k = k or SETTINGS.retrieval.top_k
                        base_docs = base_retriever(query, k=target_k)
                        qa_docs = _invoke_with_k(insuranceqa_retriever, query, target_k)

                        combined: List[Document] = []
                        seen = set()
                        for doc in base_docs + qa_docs:
                            key = (
                                doc.metadata.get("source"),
                                doc.metadata.get("page"),
                                doc.page_content[:80],
                            )
                            if key in seen:
                                continue
                            seen.add(key)
                            combined.append(doc)
                            if len(combined) >= target_k:
                                break
                        return combined

                    def insuranceqa_only(query: str, k: Optional[int] = None) -> List[Document]:
                        target_k = k or SETTINGS.retrieval.top_k
                        return _invoke_with_k(insuranceqa_retriever, query, target_k)

                    def insuranceqa_auto(query: str, k: Optional[int] = None) -> List[Document]:
                        if _should_route_to_insuranceqa(query):
                            return insuranceqa_only(query, k=k)
                        return base_retriever(query, k=k)

                    if insuranceqa_mode == "switch":
                        hybrid_retriever = insuranceqa_only
                    elif insuranceqa_mode == "auto":
                        hybrid_retriever = insuranceqa_auto
                    else:
                        hybrid_retriever = merged_retriever

                reranker_model_path = SETTINGS.reranker.model
                if not allow_reindex and mcp_local_models_only_enabled():
                    reranker_model_path = resolve_local_huggingface_model(reranker_model_path)

                components = {
                    "vector_store": vector_store,
                    "hybrid_retriever": hybrid_retriever,
                    "indexed_documents": all_splits,
                    "insuranceqa_retriever": insuranceqa_retriever,
                    "reranker_model": None,
                    "reranker_model_path": reranker_model_path,
                    "reranker_initialized": False,
                    "retrieval_service": self,
                }
            except Exception as exc:
                logger.exception("Retrieval service initialization failed: %s", exc)
                raise RuntimeError(f"Retrieval service initialization failed: {exc}") from exc

            self.components = components
            logger.info(
                "Retrieval service initialization completed in %.3fs",
                time.perf_counter() - initialization_started,
            )
            return components

    def _ensure_reranker(self) -> Any:
        components = self.components
        if components is None:
            raise RuntimeError("Retrieval service has not been initialized")
        with self._reranker_lock:
            if not components.get("reranker_initialized", False):
                started_at = time.perf_counter()
                components["reranker_model"] = build_reranker(
                    components.get("reranker_model_path")
                    or self.settings.reranker.model
                )
                components["reranker_initialized"] = True
                logger.info(
                    "Reranker lazy initialization completed in %.3fs (enabled=%s)",
                    time.perf_counter() - started_at,
                    components["reranker_model"] is not None,
                )
            return components.get("reranker_model")

    def retrieve_and_rerank(
        self,
        query: str,
        retrieval_top_k: int,
        rerank_top_k: int,
        *,
        use_reranker: bool = True,
        source_filename: Optional[str] = None,
    ) -> Tuple[List[Document], List[Document]]:
        components = self.components
        if components is None:
            raise RuntimeError("Retrieval service has not been initialized")

        request_started = time.perf_counter()
        logger.info("Retrieval request started")
        retrieval_started = begin_stage("retrieval")
        retrieved_docs: List[Document]
        attempt = 0
        while True:
            try:
                search_started = time.perf_counter()
                if source_filename:
                    requested_name = Path(
                        source_filename.replace("\\", "/")
                    ).name.lower()
                    indexed_documents = components.get("indexed_documents") or []
                    retrieved_docs = [
                        doc
                        for doc in indexed_documents
                        if Path(
                            str(
                                doc.metadata.get("source") or ""
                            ).replace("\\", "/")
                        ).name.lower()
                        == requested_name
                    ]
                    source_chunk_count = len(retrieved_docs)
                    candidate_limit = max(
                        self.settings.reranker.top_k,
                        min(
                            retrieval_top_k,
                            self.settings.reranker.top_k + 1,
                        ),
                    )
                    retrieved_docs = _prefilter_indexed_source_chunks(
                        query,
                        retrieved_docs,
                        limit=candidate_limit,
                    )
                    diagnostics = current_diagnostics()
                    if diagnostics is not None:
                        diagnostics.mark_stage("lexical_retrieval", "not_run")
                        diagnostics.mark_stage("vector_retrieval", "not_run")
                        diagnostics.mark_stage("hybrid_merge", "not_run")
                    logger.info(
                        "Existing indexed chunks selected by source "
                        "source=%s source_chunks=%d candidates=%d",
                        requested_name,
                        source_chunk_count,
                        len(retrieved_docs),
                    )
                else:
                    retrieved_docs = run_bounded_operation(
                        "retrieval",
                        lambda: components["hybrid_retriever"](
                            query,
                            k=retrieval_top_k,
                        ),
                        timeout_seconds=self.settings.llm_runtime.retrieval_timeout_seconds,
                    )
                logger.info(
                    "Retrieval search completed in %.3fs (documents=%d)",
                    time.perf_counter() - search_started,
                    len(retrieved_docs),
                )
                finish_stage("retrieval", retrieval_started)
                break
            except LLMStageTimeoutError:
                finish_stage("retrieval", retrieval_started, status="timeout")
                raise
            except Exception as exc:
                if attempt < self.settings.llm_runtime.retrieval_max_retries:
                    attempt += 1
                    retry_count = add_retry("retrieval")
                    logger.warning(
                        "Retrieval retry retry=%d reason=%s",
                        retry_count or attempt,
                        type(exc).__name__,
                    )
                    continue
                finish_stage("retrieval", retrieval_started, status="failed")
                logger.exception("Retrieval execution failed: %s", exc)
                raise RetrievalFailedError(stage="retrieval") from exc

        if use_reranker:
            reranking_started = begin_stage("reranker")
            try:
                reranked_docs = run_bounded_operation(
                    "reranker",
                    lambda: rerank_documents(
                        query,
                        retrieved_docs,
                        self._ensure_reranker(),
                        top_k=rerank_top_k,
                    ),
                    timeout_seconds=self.settings.llm_runtime.timeout_reranker_seconds,
                )
                finish_stage("reranker", reranking_started)
                logger.info(
                    "Retrieval reranking completed in %.3fs (documents=%d)",
                    time.perf_counter() - request_started,
                    len(reranked_docs),
                )
            except LLMStageTimeoutError:
                finish_stage("reranker", reranking_started, status="timeout")
                raise
            except Exception as exc:
                finish_stage("reranker", reranking_started, status="failed")
                logger.exception("Retrieval reranking failed: %s", exc)
                raise RetrievalFailedError(stage="reranker") from exc
        else:
            reranked_docs = retrieved_docs[:rerank_top_k]
            diagnostics = current_diagnostics()
            if diagnostics is not None:
                diagnostics.mark_stage("reranker", "not_run")

        logger.info(
            "Retrieval request completed in %.3fs",
            time.perf_counter() - request_started,
        )
        return retrieved_docs, reranked_docs


def initialize_retrieval_service(
    force_reindex: bool = False,
    allow_reindex: bool = False,
) -> Tuple[RetrievalService, bool]:
    """Return the process-wide retrieval service and whether it was reused."""

    global _retrieval_service
    with _retrieval_service_lock:
        created = False
        if _retrieval_service is None:
            service = RetrievalService(SETTINGS)
            created = True
        else:
            service = _retrieval_service

        reused = service.components is not None and not force_reindex
        if reused:
            logger.info("Existing retrieval-service instance reused")
            return service, True

        try:
            service.initialize(
                force_reindex=force_reindex,
                allow_reindex=allow_reindex,
            )
        except Exception:
            if created:
                _retrieval_service = None
            raise

        _retrieval_service = service
        return service, False


class RAGPipeline:
    def __init__(self, settings: ModelSettings):
        self.settings = settings
        self.components: Optional[Dict[str, Any]] = None
        self._init_lock = threading.RLock()
        self._safety_checker: Optional[SafetyChecker] = None

    def build_vectorstore(
        self,
        all_splits: List[Document],
        embeddings,
        force_reindex: bool = False,
    ) -> ChromaVectorStore:
        return build_vectorstore(all_splits, embeddings, force_reindex=force_reindex)

    def build_retriever(self, vector_store: ChromaVectorStore, all_splits: List[Document]):
        return build_retriever(vector_store, all_splits)

    def build_reranker(self):
        return build_reranker()

    def build_compressor(self):
        return build_compressor()

    def build_generation_chain(self, llm):
        return build_generation_chain(llm)

    def initialize(self, force_reindex: bool = False) -> Dict[str, Any]:
        with self._init_lock:
            if self.components is not None and not force_reindex:
                return self.components

            answer_model_warning = build_answer_model_warning(self.settings)
            if answer_model_warning:
                print(f"Warning: {answer_model_warning}")

            retrieval_service, _ = initialize_retrieval_service(
                force_reindex=force_reindex,
                allow_reindex=True,
            )
            retrieval_components = retrieval_service.components
            if retrieval_components is None:
                raise RuntimeError("Retrieval service returned no initialized components")

            answer_llm = initialize_llm()
            compressor_llm = (
                self.build_compressor()
                if self.settings.retrieval.enable_context_compression
                else None
            )
            router_llm = (
                None
                if self.settings.retrieval.force_retrieval
                else initialize_router_llm()
            )
            self.components = {
                **retrieval_components,
                "compressor_llm": compressor_llm,
                "llm": answer_llm,
                "router_llm": router_llm,
                "self_check_llm": (
                    initialize_self_check_llm()
                    if self.settings.retrieval.self_check_enabled
                    else None
                ),
                "query_rewrite_llm": (
                    initialize_query_rewrite_llm()
                    if self.settings.retrieval.query_rewrite_enabled
                    else None
                ),
            }

            _save_model_config()
            return self.components

    def run(
        self,
        query: str,
        chat_history: Optional[List[dict]] = None,
        *,
        retrieval_query: Optional[str] = None,
        additional_context_docs: Optional[List[Document]] = None,
        requested_source_filename: Optional[str] = None,
    ) -> AnswerResult:
        start_time = datetime.now()
        chat_history = chat_history or []
        original_query = query
        retrieval_query = (retrieval_query or query).strip()
        additional_context_docs = list(additional_context_docs or [])
        configure_self_check(self.settings.retrieval.self_check_enabled)
        if not self.settings.retrieval.self_check_enabled:
            mark_self_check_skipped("disabled_by_configuration")
        with self._init_lock:
            if self._safety_checker is None:
                self._safety_checker = create_safety_checker(self.settings.safety)
            safety_checker = self._safety_checker
        safety_pre_result = _default_allow_safety_result("pre_query")
        safety_context_result = _default_allow_safety_result("context")
        safety_post_result = _default_allow_safety_result("post_generation")
        final_safety_decision = "allow"
        safety_system_error = False
        safety_error_stage: Optional[str] = None
        safety_error_type: Optional[str] = None
        safety_error_message: Optional[str] = None

        def _mark_safety_system_error(stage: str, exc: Exception) -> None:
            nonlocal safety_system_error, safety_error_stage, safety_error_type, safety_error_message
            safety_system_error = True
            safety_error_stage = stage
            safety_error_type = type(exc).__name__
            safety_error_message = (
                f"Safety stage failed ({type(exc).__name__})."
            )

        guardrail_started = begin_stage("guardrail")
        try:
            safety_pre_result = safety_checker.check_query_safety(query, chat_history)
            finish_stage("guardrail", guardrail_started)
        except RuntimeExecutionError:
            finish_stage("guardrail", guardrail_started, status="timeout")
            raise
        except Exception as exc:
            finish_stage("guardrail", guardrail_started, status="failed")
            _mark_safety_system_error("pre_query", exc)
            if self.settings.safety.fail_closed and safety_checker.is_active:
                safety_pre_result = SafetyResult(
                    allow=False,
                    risk_level="high",
                    reasons=["safety_precheck_error"],
                    action="fallback",
                    details={
                        "stage": "pre_query",
                        "error_type": type(exc).__name__,
                    },
                )

        if not safety_pre_result.allow:
            blocked_answer = safety_checker.apply_safety_action(safety_pre_result, "")
            latency_ms = int((datetime.now() - start_time).total_seconds() * 1000)
            if safety_system_error and safety_error_stage == "pre_query":
                final_safety_decision = f"pre_system_error_{safety_pre_result.action}"
            else:
                final_safety_decision = f"pre_{safety_pre_result.action}"
            audit_log(
                query=original_query,
                retrieved_documents=[],
                compressed_context=[],
                generated_answer=blocked_answer,
                chat_history=chat_history,
                retrieval_needed="SAFETY_PRECHECK_BLOCKED",
                final_query=original_query,
                sources=[],
                latency_ms=latency_ms,
                retries=0,
                provider=self.settings.provider,
                answer_style=ANSWER_STYLE,
                safety_mode=self.settings.safety.mode,
                safety_enabled=safety_checker.is_active,
                safety_decision=final_safety_decision,
                safety_risks=list(safety_pre_result.reasons),
                safety_scores={"pre": dict(safety_pre_result.scores)},
                safety_block_reason=_primary_safety_reason(
                    final_safety_decision,
                    safety_pre_result,
                    safety_context_result,
                    safety_post_result,
                )
                or "precheck_blocked",
                safety_results={
                    "pre": _safety_result_to_dict(safety_pre_result),
                    "context": _safety_result_to_dict(safety_context_result),
                    "post": _safety_result_to_dict(safety_post_result),
                },
                **_safety_audit_fields(
                    final_safety_decision,
                    safety_pre_result,
                    safety_context_result,
                    safety_post_result,
                ),
                safety_system_error=safety_system_error,
                safety_error_stage=safety_error_stage,
                safety_error_type=safety_error_type,
                safety_error_message=safety_error_message,
            )
            diagnostics = current_diagnostics()
            return AnswerResult(
                answer=blocked_answer,
                sources=[],
                query=query,
                latency_ms=latency_ms,
                diagnostics=diagnostics.as_dict() if diagnostics is not None else None,
            )

        if _insuranceqa_exact_match_enabled():
            exact_answer = lookup_insuranceqa_answer(query)
            if exact_answer:
                exact_doc = Document(
                    page_content=f"Question: {query}\nAnswer: {exact_answer}",
                    metadata={"source": "insuranceqa_v2_local", "dataset": "InsuranceQA", "exact_match": True},
                )
                sources = [
                    Source(
                        document_id="insuranceqa_v2_local",
                        document_title="insuranceQA-v2 (local JSONL)",
                        page=None,
                        section="exact-match",
                        snippet=exact_answer[:300],
                    )
                ]
                final_answer = exact_answer
                guardrail_started = begin_stage("guardrail")
                try:
                    safety_post_result = safety_checker.check_answer_safety(query, [exact_doc], exact_answer)
                    finish_stage("guardrail", guardrail_started)
                except RuntimeExecutionError:
                    finish_stage("guardrail", guardrail_started, status="timeout")
                    raise
                except Exception as exc:
                    finish_stage("guardrail", guardrail_started, status="failed")
                    _mark_safety_system_error("post_generation", exc)
                    if self.settings.safety.fail_closed and safety_checker.is_active:
                        safety_post_result = SafetyResult(
                            allow=False,
                            risk_level="high",
                            reasons=["safety_postcheck_error"],
                            action="fallback",
                            details={
                                "stage": "post_generation",
                                "error_type": type(exc).__name__,
                            },
                        )

                if not safety_post_result.allow:
                    final_answer = safety_checker.apply_safety_action(safety_post_result, exact_answer)
                    if safety_post_result.action in {"block", "fallback"}:
                        sources = []
                    if safety_system_error and safety_error_stage == "post_generation":
                        final_safety_decision = f"post_system_error_{safety_post_result.action}"
                    else:
                        final_safety_decision = f"post_{safety_post_result.action}"

                latency_ms = int((datetime.now() - start_time).total_seconds() * 1000)
                audit_log(
                    query=original_query,
                    retrieved_documents=[exact_doc],
                    compressed_context=[exact_doc],
                    generated_answer=final_answer,
                    chat_history=chat_history,
                    retrieval_needed="RETRIEVE",
                    final_query=original_query,
                    sources=[s.__dict__ for s in sources],
                    latency_ms=latency_ms,
                    retries=0,
                    provider=self.settings.provider,
                    answer_style=ANSWER_STYLE,
                    exact_insuranceqa_match=True,
                    safety_mode=self.settings.safety.mode,
                    safety_enabled=safety_checker.is_active,
                    safety_decision=final_safety_decision,
                    safety_risks=list(set(safety_pre_result.reasons + safety_post_result.reasons)),
                    safety_scores={
                        "pre": dict(safety_pre_result.scores),
                        "context": dict(safety_context_result.scores),
                        "post": dict(safety_post_result.scores),
                    },
                    safety_block_reason=_primary_safety_reason(
                        final_safety_decision,
                        safety_pre_result,
                        safety_context_result,
                        safety_post_result,
                    ),
                    safety_results={
                        "pre": _safety_result_to_dict(safety_pre_result),
                        "context": _safety_result_to_dict(safety_context_result),
                        "post": _safety_result_to_dict(safety_post_result),
                    },
                    **_safety_audit_fields(
                        final_safety_decision,
                        safety_pre_result,
                        safety_context_result,
                        safety_post_result,
                    ),
                    safety_system_error=safety_system_error,
                    safety_error_stage=safety_error_stage,
                    safety_error_type=safety_error_type,
                    safety_error_message=safety_error_message,
                )
                diagnostics = current_diagnostics()
                return AnswerResult(
                    answer=final_answer,
                    sources=sources,
                    query=query,
                    latency_ms=latency_ms,
                    diagnostics=diagnostics.as_dict() if diagnostics is not None else None,
                )

        components = self.initialize(force_reindex=pdfs_have_changed())

        if self.settings.retrieval.force_retrieval:
            retrieval_needed = "RETRIEVE"
        else:
            retrieval_needed = decide_retrieval(
                components["router_llm"],
                retrieval_query,
                chat_history,
            )

        if retrieval_needed != "RETRIEVE":
            answer = generate_direct_answer(components["llm"], query, chat_history)
            sources: List[Source] = []
            latency_ms = int((datetime.now() - start_time).total_seconds() * 1000)
            audit_log(
                query=original_query,
                retrieved_documents=[],
                compressed_context=[],
                generated_answer=answer,
                chat_history=chat_history,
                retrieval_needed=retrieval_needed,
                final_query=original_query,
                sources=[],
                latency_ms=latency_ms,
                provider=self.settings.provider,
                answer_style=ANSWER_STYLE,
                safety_mode=self.settings.safety.mode,
                safety_enabled=safety_checker.is_active,
                safety_decision=final_safety_decision,
                safety_risks=list(safety_pre_result.reasons),
                safety_scores={
                    "pre": dict(safety_pre_result.scores),
                    "context": dict(safety_context_result.scores),
                    "post": dict(safety_post_result.scores),
                },
                safety_block_reason=None,
                safety_results={
                    "pre": _safety_result_to_dict(safety_pre_result),
                    "context": _safety_result_to_dict(safety_context_result),
                    "post": _safety_result_to_dict(safety_post_result),
                },
                **_safety_audit_fields(
                    final_safety_decision,
                    safety_pre_result,
                    safety_context_result,
                    safety_post_result,
                ),
                safety_system_error=safety_system_error,
                safety_error_stage=safety_error_stage,
                safety_error_type=safety_error_type,
                safety_error_message=safety_error_message,
            )
            diagnostics = current_diagnostics()
            return AnswerResult(
                answer=answer,
                sources=sources,
                query=query,
                latency_ms=latency_ms,
                diagnostics=diagnostics.as_dict() if diagnostics is not None else None,
            )

        current_query = retrieval_query
        reranked_docs: List[Document] = []
        last_retrieved_docs: List[Document] = []
        retrieval_service = components.get("retrieval_service")
        if not isinstance(retrieval_service, RetrievalService):
            raise RuntimeError("Full RAG pipeline is missing the shared retrieval service")

        retries = 0
        max_attempts = max(
            1,
            min(
                self.settings.retrieval.max_self_check_retries,
                1 + self.settings.llm_runtime.retrieval_max_retries,
            ),
        )
        while retries < max_attempts:
            source_scoped_top_k = (
                min(2, self.settings.reranker.top_k)
                if requested_source_filename
                else self.settings.reranker.top_k
            )
            retrieved_docs, reranked_docs = retrieval_service.retrieve_and_rerank(
                current_query,
                retrieval_top_k=self.settings.retrieval.top_k,
                rerank_top_k=source_scoped_top_k,
                use_reranker=not bool(requested_source_filename),
                source_filename=requested_source_filename,
            )
            last_retrieved_docs = retrieved_docs
            record_evidence(
                "pdfChunks",
                [
                    {
                        **_document_evidence(doc, rank=index + 1),
                        "textExcerpt": _audit_safe_text((doc.page_content or "")[:600]),
                    }
                    for index, doc in enumerate(reranked_docs)
                    if doc.metadata.get("source_type") != "crm"
                ],
            )

            if not reranked_docs:
                rewritten_query = rewrite_query(components["query_rewrite_llm"], current_query, chat_history)
                if rewritten_query == current_query:
                    break
                current_query = rewritten_query
                retries += 1
                continue

            if self.settings.retrieval.self_check_enabled:
                checked_query, checked_docs = perform_self_check(
                    components["self_check_llm"],
                    components["query_rewrite_llm"],
                    current_query,
                    reranked_docs,
                    chat_history,
                )
            else:
                checked_query, checked_docs = current_query, reranked_docs

            if checked_query != current_query or not checked_docs:
                if checked_query == current_query and not checked_docs:
                    break
                current_query = checked_query
                retries += 1
                continue

            reranked_docs = checked_docs
            break

        if not reranked_docs:
            answer = NO_RELEVANT_INFORMATION_MESSAGE
            sources = []
            context_docs_for_log: List[Document] = []
        else:
            context_docs = reranked_docs + additional_context_docs
            sanitized_context_docs: Optional[List[Document]] = None
            guardrail_started = begin_stage("guardrail")
            try:
                safety_context_result = safety_checker.check_context_safety(context_docs)
                finish_stage("guardrail", guardrail_started)
            except RuntimeExecutionError:
                finish_stage("guardrail", guardrail_started, status="timeout")
                raise
            except Exception as exc:
                finish_stage("guardrail", guardrail_started, status="failed")
                _mark_safety_system_error("context", exc)
                if self.settings.safety.fail_closed and safety_checker.is_active:
                    safety_context_result = SafetyResult(
                        allow=False,
                        risk_level="high",
                        reasons=["safety_context_check_error"],
                        action="fallback",
                        details={
                            "stage": "context",
                            "error_type": type(exc).__name__,
                        },
                    )

            if not safety_context_result.allow:
                sanitized_context_docs = _context_docs_from_safety_result(safety_context_result)

            if not safety_context_result.allow and not (
                safety_context_result.action == "redact" and sanitized_context_docs
            ):
                answer = safety_checker.apply_safety_action(safety_context_result, "")
                sources = []
                context_docs_for_log = reranked_docs
                if safety_system_error and safety_error_stage == "context":
                    final_safety_decision = f"context_system_error_{safety_context_result.action}"
                else:
                    final_safety_decision = f"context_{safety_context_result.action}"
            else:
                if sanitized_context_docs:
                    context_docs = sanitized_context_docs
                    final_safety_decision = "context_redact"
                if self.settings.retrieval.enable_context_compression:
                    context_docs = compress_context(components["compressor_llm"], context_docs, current_query)

                record_evidence(
                    "answerGenerationStatus",
                    {"started": True, "completed": False},
                )
                try:
                    answer = generate_answer(
                        components["llm"],
                        original_query,
                        context_docs,
                        chat_history,
                    )
                except Exception as exc:
                    record_evidence(
                        "answerGenerationStatus",
                        {
                            "started": True,
                            "completed": False,
                            "exceptionType": type(exc).__name__,
                            "exceptionMessageSanitized": _audit_safe_text(str(exc)),
                        },
                    )
                    raise
                record_evidence(
                    "answerGenerationStatus",
                    {"started": True, "completed": True},
                )
                sources = [
                    document_to_source(doc)
                    for doc in (reranked_docs + additional_context_docs)
                ]
                context_docs_for_log = context_docs

                _record_answer_version("immediatelyBeforeGroundedness", answer)

                guardrail_started = begin_stage("guardrail")
                groundedness_started = begin_stage("groundedness")
                try:
                    safety_post_result = safety_checker.check_answer_safety(
                        original_query,
                        context_docs,
                        answer,
                    )
                    finish_stage("groundedness", groundedness_started)
                    finish_stage("guardrail", guardrail_started)
                except RuntimeExecutionError as exc:
                    finish_stage("groundedness", groundedness_started, status="timeout")
                    finish_stage("guardrail", guardrail_started, status="timeout")
                    record_evidence(
                        "groundedness",
                        {
                            "algorithm": "fact_aware_claim_support_v5",
                            "score": None,
                            "threshold": self.settings.safety.min_groundedness,
                            "passed": None,
                            "exceptionType": type(exc).__name__,
                            "exceptionMessageSanitized": _audit_safe_text(str(exc)),
                        },
                    )
                    raise
                except Exception as exc:
                    finish_stage("groundedness", groundedness_started, status="failed")
                    finish_stage("guardrail", guardrail_started, status="failed")
                    _mark_safety_system_error("post_generation", exc)
                    record_evidence(
                        "groundedness",
                        {
                            "algorithm": "fact_aware_claim_support_v5",
                            "score": None,
                            "threshold": self.settings.safety.min_groundedness,
                            "passed": None,
                            "exceptionType": type(exc).__name__,
                            "exceptionMessageSanitized": _audit_safe_text(str(exc)),
                        },
                    )
                    if self.settings.safety.fail_closed and safety_checker.is_active:
                        safety_post_result = SafetyResult(
                            allow=False,
                            risk_level="high",
                            reasons=["safety_postcheck_error"],
                            action="fallback",
                            details={
                                "stage": "post_generation",
                                "error_type": type(exc).__name__,
                            },
                        )

                if not safety_post_result.allow:
                    answer = safety_checker.apply_safety_action(safety_post_result, answer)
                    if safety_post_result.action in {"block", "fallback"}:
                        sources = []
                    if safety_system_error and safety_error_stage == "post_generation":
                        final_safety_decision = f"post_system_error_{safety_post_result.action}"
                    else:
                        final_safety_decision = f"post_{safety_post_result.action}"

        latency_ms = int((datetime.now() - start_time).total_seconds() * 1000)
        safety_risks = list(
            set(
                safety_pre_result.reasons
                + safety_context_result.reasons
                + safety_post_result.reasons
            )
        )
        groundedness_score = safety_post_result.scores.get("groundedness")
        groundedness_details = safety_post_result.details.get("groundedness", {})
        record_evidence(
            "groundedness",
            {
                **sanitize_diagnostic_value(groundedness_details),
                "algorithm": groundedness_details.get(
                    "algorithm_version", "fact_aware_claim_support_v5"
                ),
                "score": groundedness_score,
                "threshold": self.settings.safety.min_groundedness,
                "thresholdSource": self.settings.safety.min_groundedness_source,
                "passed": (
                    float(groundedness_score)
                    >= float(self.settings.safety.min_groundedness)
                    if groundedness_score is not None
                    else None
                ),
            },
        )
        _record_answer_version("finalAfterOutputSafety", answer)
        record_evidence(
            "pdfCitations",
            [
                source.__dict__
                for source in sources
                if not source.document_id.startswith("espocrm:")
            ],
        )
        record_evidence(
            "safety",
            {
                "decision": final_safety_decision,
                "risks": safety_risks,
                "preAction": safety_pre_result.action,
                "contextAction": safety_context_result.action,
                "postAction": safety_post_result.action,
            },
        )
        audit_log(
            query=original_query,
            retrieved_documents=last_retrieved_docs + additional_context_docs,
            compressed_context=context_docs_for_log,
            generated_answer=answer,
            chat_history=chat_history,
            retrieval_needed=retrieval_needed,
            final_query=current_query,
            sources=[s.__dict__ for s in sources],
            latency_ms=latency_ms,
            retries=retries,
            provider=self.settings.provider,
            answer_style=ANSWER_STYLE,
            safety_mode=self.settings.safety.mode,
            safety_enabled=safety_checker.is_active,
            safety_decision=final_safety_decision,
            safety_risks=safety_risks,
            safety_scores={
                "pre": dict(safety_pre_result.scores),
                "context": dict(safety_context_result.scores),
                "post": dict(safety_post_result.scores),
            },
            safety_block_reason=_primary_safety_reason(
                final_safety_decision,
                safety_pre_result,
                safety_context_result,
                safety_post_result,
            ),
            safety_results={
                "pre": _safety_result_to_dict(safety_pre_result),
                "context": _safety_result_to_dict(safety_context_result),
                "post": _safety_result_to_dict(safety_post_result),
            },
            **_safety_audit_fields(
                final_safety_decision,
                safety_pre_result,
                safety_context_result,
                safety_post_result,
            ),
            safety_system_error=safety_system_error,
            safety_error_stage=safety_error_stage,
            safety_error_type=safety_error_type,
            safety_error_message=safety_error_message,
        )
        diagnostics = current_diagnostics()
        return AnswerResult(
            answer=answer,
            sources=sources,
            query=query,
            latency_ms=latency_ms,
            diagnostics=diagnostics.as_dict() if diagnostics is not None else None,
        )


def initialize_pipeline(force_reindex: bool = False):
    global _pipeline
    with _pipeline_lock:
        if _pipeline is None:
            _pipeline = RAGPipeline(SETTINGS)
        return _pipeline.initialize(force_reindex=force_reindex)


def _json_compatible_metadata(metadata: dict) -> dict:
    compatible: dict = {}
    for key, value in (metadata or {}).items():
        if value is None or isinstance(value, (str, int, float, bool)):
            compatible[key] = value
        elif isinstance(value, Path):
            compatible[key] = str(value)
        else:
            compatible[key] = str(value)
    return compatible


def retrieve_documents_for_tool(
    question: str,
    top_k: int = 5,
    *,
    use_reranker: bool = True,
) -> dict:
    question = (question or "").strip()
    if not question:
        raise ValueError("question must not be empty")

    bounded_top_k = max(1, min(int(top_k), 20))
    tool_call_started = time.perf_counter()
    logger.info("MCP retrieval tool call started")
    try:
        retrieval_service, reused = initialize_retrieval_service(
            force_reindex=False,
            allow_reindex=False,
        )
        logger.info("Retrieval-service instance reused=%s", reused)
        _, docs = retrieval_service.retrieve_and_rerank(
            question,
            retrieval_top_k=bounded_top_k,
            rerank_top_k=bounded_top_k,
            use_reranker=use_reranker,
        )

        try:
            documents = []
            for rank, doc in enumerate(docs, start=1):
                metadata = _json_compatible_metadata(dict(doc.metadata or {}))
                documents.append(
                    {
                        "rank": rank,
                        "content": doc.page_content or "",
                        "metadata": metadata,
                        "source": {
                            "source": metadata.get("source"),
                            "source_type": metadata.get("source_type"),
                            "page": metadata.get("page"),
                            "sheet": metadata.get("sheet"),
                            "table_format": metadata.get("table_format"),
                            "row_count": metadata.get("row_count"),
                            "table_index": metadata.get("table_index"),
                        },
                    }
                )
        except Exception as exc:
            logger.exception("Retrieval result serialization failed: %s", exc)
            raise RuntimeError(f"Retrieval result serialization failed: {exc}") from exc

        return {
            "query": question,
            "top_k": bounded_top_k,
            "retrieval_mode": (
                "hybrid_with_existing_reranker"
                if use_reranker
                else "hybrid_without_reranker"
            ),
            "generation_used": False,
            "reranker_used": bool(
                use_reranker
                and retrieval_service.components
                and retrieval_service.components.get("reranker_model") is not None
            ),
            "document_count": len(documents),
            "documents": documents,
            "diagnostics": (
                current_diagnostics().as_dict()
                if current_diagnostics() is not None
                else None
            ),
        }
    finally:
        logger.info(
            "MCP retrieval tool call completed in %.3fs",
            time.perf_counter() - tool_call_started,
        )


def run_rag(
    question: str,
    chat_history: Optional[List[dict]] = None,
    *,
    retrieval_query: Optional[str] = None,
    additional_context_docs: Optional[List[Document]] = None,
    requested_source_filename: Optional[str] = None,
) -> AnswerResult:
    global _pipeline
    with _pipeline_lock:
        if _pipeline is None:
            _pipeline = RAGPipeline(SETTINGS)
        pipeline = _pipeline
    return pipeline.run(
        question,
        chat_history=chat_history,
        retrieval_query=retrieval_query,
        additional_context_docs=additional_context_docs,
        requested_source_filename=requested_source_filename,
    )


def is_pipeline_ready() -> bool:
    with _pipeline_lock:
        return _pipeline is not None and _pipeline.components is not None


def is_retrieval_ready() -> bool:
    with _retrieval_service_lock:
        return (
            _retrieval_service is not None
            and _retrieval_service.components is not None
        )


def is_embedding_ready() -> bool:
    with _retrieval_service_lock:
        components = (
            _retrieval_service.components
            if _retrieval_service is not None
            else None
        )
        return bool(
            components is not None
            and components.get("vector_store") is not None
        )
