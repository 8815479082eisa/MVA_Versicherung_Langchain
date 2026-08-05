from __future__ import annotations

import argparse
import csv
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from sklearn.metrics import cohen_kappa_score


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_INPUT = PROJECT_ROOT / "reports" / "groundedness_dataset_review.csv"
DEFAULT_JSON = PROJECT_ROOT / "reports" / "groundedness_annotation_agreement.json"
DEFAULT_CONFLICTS = PROJECT_ROOT / "reports" / "groundedness_annotation_conflicts.csv"
LABELS = {"PASS", "FAIL", "AMBIGUOUS", "EXCLUDE"}
CRITICALITIES = {"CRITICAL", "NON_CRITICAL", "REVIEW_REQUIRED"}


def _clean(value: Any) -> str:
    return "" if value is None else str(value).strip().upper()


def _read_csv(path: Path) -> list[dict[str, Any]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def _read_xlsx(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("Reading XLSX requires the project's existing openpyxl dependency") from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = [name for name in ("Reviewer_1", "Reviewer_2", "Adjudication") if name in workbook.sheetnames]
    if not sheets:
        raise ValueError("Workbook must contain Reviewer_1, Reviewer_2, or Adjudication sheets")
    merged: dict[str, dict[str, Any]] = {}
    for sheet_name in sheets:
        sheet = workbook[sheet_name]
        rows = sheet.iter_rows(values_only=True)
        headers = [str(value or "").strip() for value in next(rows)]
        for values in rows:
            row = dict(zip(headers, values))
            case_id = str(row.get("case_id") or "").strip()
            if not case_id:
                continue
            target = merged.setdefault(case_id, {"case_id": case_id})
            for key, value in row.items():
                if key and value not in (None, ""):
                    target[key] = value
    return list(merged.values())


def read_rows(path: Path) -> list[dict[str, Any]]:
    return _read_xlsx(path) if path.suffix.lower() == ".xlsx" else _read_csv(path)


def _agreement(rows: list[dict[str, Any]], left: str, right: str, allowed: set[str]) -> dict[str, Any]:
    paired = [(_clean(row.get(left)), _clean(row.get(right))) for row in rows]
    paired = [(a, b) for a, b in paired if a in allowed and b in allowed]
    agreed = sum(a == b for a, b in paired)
    kappa = float(cohen_kappa_score([a for a, _ in paired], [b for _, b in paired])) if len(paired) >= 2 else None
    if kappa is not None and (kappa != kappa):
        kappa = None
    return {
        "paired_cases": len(paired),
        "agreement_count": agreed,
        "raw_agreement": agreed / len(paired) if paired else None,
        "cohen_kappa": kappa,
    }


def analyze(rows: list[dict[str, Any]]) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    ids = [str(row.get("case_id") or "").strip() for row in rows]
    if not rows or any(not case_id for case_id in ids) or len(ids) != len(set(ids)):
        raise ValueError("Annotation input must contain unique, non-empty case_id values")

    conflicts: list[dict[str, Any]] = []
    for row in rows:
        r1 = _clean(row.get("reviewer_1_label"))
        r2 = _clean(row.get("reviewer_2_label"))
        category_1 = _clean(row.get("reviewer_1_error_category"))
        category_2 = _clean(row.get("reviewer_2_error_category"))
        criticality_1 = _clean(row.get("reviewer_1_criticality"))
        criticality_2 = _clean(row.get("reviewer_2_criticality"))
        reasons: list[str] = []
        if r1 in LABELS and r2 in LABELS and r1 != r2:
            reasons.append("label")
        if category_1 and category_2 and category_1 != category_2:
            reasons.append("error_category")
        if criticality_1 in CRITICALITIES and criticality_2 in CRITICALITIES and criticality_1 != criticality_2:
            reasons.append("criticality")
        if reasons:
            conflicts.append(
                {
                    "case_id": row["case_id"],
                    "conflict_fields": ";".join(reasons),
                    "reviewer_1_label": row.get("reviewer_1_label", ""),
                    "reviewer_2_label": row.get("reviewer_2_label", ""),
                    "reviewer_1_error_category": row.get("reviewer_1_error_category", ""),
                    "reviewer_2_error_category": row.get("reviewer_2_error_category", ""),
                    "reviewer_1_criticality": row.get("reviewer_1_criticality", ""),
                    "reviewer_2_criticality": row.get("reviewer_2_criticality", ""),
                    "adjudicated_label": row.get("adjudicated_label", ""),
                    "adjudication_notes": row.get("adjudication_notes", ""),
                }
            )

    reviewer_1_complete = sum(_clean(row.get("reviewer_1_label")) in LABELS for row in rows)
    reviewer_2_complete = sum(_clean(row.get("reviewer_2_label")) in LABELS for row in rows)
    adjudicated_complete = sum(_clean(row.get("adjudicated_label")) in LABELS for row in rows)
    unresolved_conflicts = sum(
        not _clean(row.get("adjudicated_label")) in LABELS or not str(row.get("adjudication_notes") or "").strip()
        for row in conflicts
    )
    result = {
        "schema_version": 1,
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "case_count": len(rows),
        "status": "READY_FOR_ISOLATED_SCORING" if (
            reviewer_1_complete == len(rows)
            and reviewer_2_complete == len(rows)
            and adjudicated_complete == len(rows)
            and unresolved_conflicts == 0
        ) else "ANNOTATION_PENDING",
        "completion": {
            "reviewer_1": reviewer_1_complete,
            "reviewer_2": reviewer_2_complete,
            "adjudicated": adjudicated_complete,
            "total": len(rows),
        },
        "label_agreement": _agreement(rows, "reviewer_1_label", "reviewer_2_label", LABELS),
        "pass_fail_agreement_and_cohen_kappa": _agreement(
            rows, "reviewer_1_label", "reviewer_2_label", {"PASS", "FAIL"}
        ),
        "criticality_agreement": _agreement(
            rows, "reviewer_1_criticality", "reviewer_2_criticality", CRITICALITIES
        ),
        "category_exact_agreement": _agreement(
            rows,
            "reviewer_1_error_category",
            "reviewer_2_error_category",
            ({
                _clean(row.get("reviewer_1_error_category")) for row in rows
            } | {
                _clean(row.get("reviewer_2_error_category")) for row in rows
            }) - {""},
        ),
        "reviewer_1_label_counts": dict(Counter(_clean(row.get("reviewer_1_label")) or "BLANK" for row in rows)),
        "reviewer_2_label_counts": dict(Counter(_clean(row.get("reviewer_2_label")) or "BLANK" for row in rows)),
        "conflict_count": len(conflicts),
        "unresolved_conflict_count": unresolved_conflicts,
        "warning": "Generator expectations and groundedness scores are not used as human labels.",
    }
    return result, conflicts


def main() -> int:
    parser = argparse.ArgumentParser(description="Analyze two independent groundedness annotations")
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-json", type=Path, default=DEFAULT_JSON)
    parser.add_argument("--conflicts-csv", type=Path, default=DEFAULT_CONFLICTS)
    args = parser.parse_args()

    result, conflicts = analyze(read_rows(args.input))
    args.output_json.parent.mkdir(parents=True, exist_ok=True)
    args.output_json.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    args.conflicts_csv.parent.mkdir(parents=True, exist_ok=True)
    fields = [
        "case_id", "conflict_fields", "reviewer_1_label", "reviewer_2_label",
        "reviewer_1_error_category", "reviewer_2_error_category",
        "reviewer_1_criticality", "reviewer_2_criticality",
        "adjudicated_label", "adjudication_notes",
    ]
    with args.conflicts_csv.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(conflicts)
    print(json.dumps(result, indent=2, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
